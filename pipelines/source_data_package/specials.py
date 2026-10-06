"""Content changes for specific package paths (applied before the generic redaction).

Each function takes the original bytes and returns (new_bytes, note)."""
import csv, io, os, re, sys

sys.path.insert(0, os.path.dirname(__file__))
import selection as S

ST4_README = [
    'Supplementary Table 5. Individual-function correlations between DNA, RNA and protein layers in the PREDICT cohort.',
    '',
    'All_30581_tests: one row per eligible function and layer pair under the revised allocation (30,581 tests): number of patients, Spearman rho, two-sided P value and Benjamini-Hochberg q value within each of the nine representation-by-layer-pair families.',
    'Nine_family_summary: for each family, the number of tested functions, the median rho and the number of functions with q < 0.05.',
    'All_eligibility: every function and layer pair considered, with the number of patients with positive measurements in both layers and whether it met the eligibility rule (at least 35 patients).',
    'Eligibility_common: for each family, the number of functions eligible under both the previous and the revised allocation (n_common), eligible only under the revised allocation (n_newly_eligible) and eligible only under the previous allocation (n_lost), with the median revised correlation for common and for newly eligible functions and the median previous correlation for common functions. Profiles were not renormalised, and correlations, P values and q values were not refitted for this subset.',
    '',
    'Previous allocation: the protein-to-function allocation used before the annotation revision of 2 October 2026. Revised allocation: the allocation used for all other values in this table.',
    'Native UHGG eggNOG annotations were used only when the protein aligned to the annotation seed orthologue with at least half of the reported bitscore (Methods). Values computed before this check are in Source Data, Annotation_verification.',
    'Per-function rows of the eligibility comparison: Source Data, Supplementary_Tables/Table5_support/eligibility_matched_correlations.tsv.',
]
ELIG_COLS = {'n_patients_refit': 'n_patients_revised_allocation', 'rho_refit': 'rho_revised_allocation',
             'n_patients_baseline': 'n_patients_previous_allocation', 'rho_baseline': 'rho_previous_allocation'}
ELIG_VALS = {'both': 'common', 'left_only': 'newly_eligible', 'right_only': 'lost'}
SUM_COLS = {'n_new': 'n_newly_eligible', 'median_refit_common': 'median_rho_revised_common',
            'median_baseline_common': 'median_rho_previous_common', 'median_refit_new': 'median_rho_revised_newly_eligible'}


def _tsv(data):
    return list(csv.reader(io.StringIO(data.decode('utf-8')), delimiter='\t'))


def _write_tsv(rows):
    out = io.StringIO(); csv.writer(out, delimiter='\t', lineterminator='\n').writerows(rows)
    return out.getvalue().encode('utf-8')


def st4(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    removed = [n for n in wb.sheetnames if n not in ('All_30581_tests', 'Nine_family_summary', 'All_eligibility', 'Eligibility_common')]
    for n in removed:
        del wb[n]
    ws = wb.create_sheet('Read_me', 0)
    for line in ST4_README:
        ws.append([line])
    ec = wb['Eligibility_common']
    for c in ec[1]:
        if c.value in SUM_COLS: c.value = SUM_COLS[c.value]
    out = io.BytesIO(); wb.save(out)
    return out.getvalue(), f'sheets {",".join(removed)} replaced by a Read_me sheet; Eligibility_common columns renamed'


ST3_README = [
    'Supplementary Table 4. Enterotype Dysbiosis Score (EDS) displays in the PREDICT cohort.',
    '',
    'All20_features_259_observations, Fixed_feature_key, Observation_metadata, 85_patient_order, 85_patient_raw and 85_patient_z: the earlier 20-feature, 259-observation selection and its 85-patient display (Supplementary Fig. 10b). This selection was fixed earlier and was not repeated. HV366 remains in these tables and has no positive signal in the patient display.',
    'Focused_raw: the six measurements displayed for the three metabolite-enzyme pairs in Fig. 4h and Supplementary Fig. 10a, with the patient order. Non-detections remain missing in the plots.',
    'Focused_features: the associations of the six displayed features. Focused_pairs: q values saved from the complete screen families, not a correction restricted to the three displayed pairs. Focused_adjustments: all 192 descriptive model estimates. No new discovery P values or intervals were computed.',
    'Human rows are human-assigned peptide signal within retained reporting groups. Pathway rows exclude human-assigned signal and retain unresolved origin.',
    'Non-detections are not imputed. This descriptive display is not an independent validation of the EDS.',
    'KO and pathway values use native UHGG eggNOG annotations only when the protein aligned to the annotation seed orthologue with at least half of the reported bitscore (Methods). Values computed before this check are in Source Data, Annotation_verification.',
    'Complete screen tables: Source Data, EDS_focused_20261003 and Biology_refit/eds_refit.',
]


def st3(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    removed = [n for n in ('Current_EDS_update', 'Read_me') if n in wb.sheetnames]
    for n in removed:
        del wb[n]
    ws = wb.create_sheet('Read_me', 0)
    for line in ST3_README:
        ws.append([line])
    out = io.BytesIO(); wb.save(out)
    return out.getvalue(), f'sheets {",".join(removed)} replaced by one neutral Read_me sheet'


def elig(data):
    rows = _tsv(data)
    h = rows[0]
    rows[0] = [ELIG_COLS.get(x, x) for x in h]
    j = h.index('eligibility')
    bad = {r[j] for r in rows[1:] if r[j] not in ELIG_VALS}
    if bad:
        raise SystemExit(f'unexpected eligibility labels {bad}')
    for r in rows[1:]:
        r[j] = ELIG_VALS[r[j]]
    return _write_tsv(rows), 'eligibility labels both/left_only relabelled common/newly_eligible; refit/baseline columns renamed revised/previous allocation'


def elig_sum(data):
    rows = _tsv(data)
    rows[0] = [SUM_COLS.get(x, x) for x in rows[0]]
    return _write_tsv(rows), 'columns renamed (newly eligible; revised/previous allocation)'


ST6_README = [
    'Supplementary Table 3. PREDICT antibiotic endpoints under alternative handling of rifaximin for hepatic encephalopathy.',
    '',
    'Rows give the primary definition (42 patients), rifaximin for hepatic encephalopathy counted as exposure (40 patients), patients with any such record excluded (35 patients) and only off-state acquisitions with such a record excluded (38 patients).',
    'Endpoints: for human-assigned share of total LFQ intensity (percentage points) and quantified protein groups, the number of paired patients and acquisitions, state means, mean and median paired differences (on minus off), number of patients with a higher value on antibiotics, 95% bootstrap interval of the mean difference, sign-flip P value and Benjamini-Hochberg q value across the two endpoints within each definition.',
    'Composition: mean origin composition of total LFQ intensity (fractions) by antibiotic state for each definition.',
    'All definitions used the same resampling settings and seeds as the primary analysis.',
]


def st6(data):
    import openpyxl
    rows = _tsv(data)
    comp_path = os.path.join(S.SOURCES['DRIF'][1], 'out/paired_composition_by_scenario.tsv')
    comp = _tsv(open(comp_path, 'rb').read())
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Read_me'
    for line in ST6_README: ws.append([line])
    for title, tab in (('Endpoints', rows), ('Composition', comp)):
        w = wb.create_sheet(title)
        for i, r in enumerate(tab):
            if i == 0:
                w.append(r); continue
            conv = []
            for x in r:
                try:
                    conv.append(int(x) if re.fullmatch(r'-?\d+', x) else float(x))
                except ValueError:
                    conv.append(x)
            w.append(conv)
    wb.properties.creator = 'FastaLake'
    out = io.BytesIO(); wb.save(out)
    return out.getvalue(), 'workbook assembled from the rifaximin sensitivity tables (endpoints and composition) with a Read_me sheet'


EGG_COLS = ['quantity', 'previous_rows_as_in_manuscript', 'verified_rows_threshold_0.5_primary', 'verified_rows_threshold_0.8_sensitivity', 'change_to_manuscript_claim']


def egg(data):
    rows = _tsv(data)
    h = rows[0]
    if h[:5] != ['quantity', 'frozen', 'corrected_verified50_primary', 'corrected_verified80_sensitivity', 'claim_changes']:
        raise SystemExit(f'unexpected header {h}')
    out = [EGG_COLS] + [r[:5] for r in rows[1:]]
    for r in out[1:]:
        r[1] = r[1].replace('taxonomy_signal frozen', 'unchanged input').replace('capscan_baseline (earlier allocation)', 'earlier allocation')
    return _write_tsv(out), 'producer file-path column removed; columns renamed'


def renum_tables(data):
    """ROUND3: the comparison tables name Supplementary Tables in the v20 numbering; move them to v21 (EDS 3 -> 4, individual functions 4 -> 5)."""
    t = data.decode('utf-8')
    t2 = re.sub(r'Supplementary([ _])Table([ _])([34])', lambda m: f'Supplementary{m.group(1)}Table{m.group(2)}' + {'3': '4', '4': '5'}[m.group(3)], t)
    t2 = re.sub(r'Table4_support', 'Table5_support', t2)
    return t2.encode('utf-8'), 'Supplementary Table numbers updated to the manuscript numbering (3 -> 4, 4 -> 5)'


SPECIAL = {
    'Annotation_verification/previous_vs_verified_all_numbers.tsv': renum_tables,
    'Annotation_verification/previous_vs_verified_printed_numbers.tsv': renum_tables,
    'Supplementary_Tables/Supplementary_Table_4.xlsx': st3,
    'Supplementary_Tables/Supplementary_Table_5.xlsx': st4,
    'Supplementary_Tables/Table5_support/eligibility_matched_correlations.tsv': elig,
    'Supplementary_Tables/Table5_support/common_eligibility_summary.tsv': elig_sum,
    'Supplementary_Tables/Supplementary_Table_3.xlsx': st6,
    'Annotation_verification/native_eggNOG_verification_previous_vs_verified.tsv': egg,
}


def apply(rel, data):
    f = SPECIAL.get(rel)
    if not f:
        return data, ''
    return f(data)
