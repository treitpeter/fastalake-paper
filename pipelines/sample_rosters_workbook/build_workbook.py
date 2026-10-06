"""Package exact analysis rosters; never harmonize source identities implicitly."""
from pathlib import Path
from collections import Counter
import csv, hashlib, json, zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path.cwd().resolve()
O = Path(__file__).resolve().parent
D = O / 'delivery'
D.mkdir(exist_ok=True)

def read(rel):
    with (O / rel).open(newline='') as f:
        return list(csv.DictReader(f, delimiter='\t'))

def digest(p):
    h = hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda: f.read(1024*1024), b''):
            h.update(b)
    return h.hexdigest()

def write_tsv(path, rows):
    keys = list(dict.fromkeys(k for r in rows for k in r))
    with path.open('w', newline='') as f:
        w = csv.DictWriter(f, keys, delimiter='\t', extrasaction='raise')
        w.writeheader(); w.writerows(rows)

wb = Workbook(); wb.remove(wb.active)
sheet_records = []
expected = {}

def sheet(name, rows, description, source='Derived by build_workbook.py'):
    assert rows, name
    ws = wb.create_sheet(name)
    keys = list(dict.fromkeys(k for r in rows for k in r))
    values = [[str(r.get(k, '')) if r.get(k, '') is not None else '' for k in keys] for r in rows]
    ws.append(keys)
    for row in values:
        ws.append(row)
    # Store identifiers as text, including numeric-looking IDs and leading zeroes.
    for row in ws:
        for c in row:
            c.data_type = 's'
            c.alignment = Alignment(vertical='top', wrap_text=True)
    for c in ws[1]:
        c.font = Font(name='Calibri', bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='24475B')
    for j, key in enumerate(keys, 1):
        maxlen = max([len(key)] + [len(r[j-1]) for r in values[:100]])
        wide = any(s in key.lower() for s in ['reason', 'rule', 'purpose', 'limit', 'note', 'interpret', 'description'])
        ws.column_dimensions[get_column_letter(j)].width = min(70 if wide else 48, max(18, maxlen + 2))
    ws.freeze_panes = 'C2' if len(keys) > 2 else 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 85
    ws.row_dimensions[1].height = 32
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    expected[name] = [keys] + values
    sheet_records.append(dict(sheet=name, data_rows=len(rows), description=description, source_table=source))
    return ws

notes = [
    ('Purpose', 'Exact measurement membership, sample/participant links and recorded exclusions for the main Methods review (M01). Start with Cohorts and Analyses; filter the named roster sheet by analysis_id.'),
    ('Units', 'An acquisition is an LC–MS/MS measurement. A specimen is biological material identified in the producing metadata. Repeated acquisitions, fractions, preparations and participant-region summaries are not additional independent people.'),
    ('Unknown values', 'Blank source identifiers remain blank. No specimen or original donor identifier has been invented. HS_shared_donor is an analysis label for the author-confirmed shared donor, not a recovered source identity.'),
    ('Healthy stool', 'Peter confirms that OnePiece, technical reinjections and workflow preparations came from one donor. The author-supplied protocol links OnePiece to longitudinal sample 2. Whether this donor belongs to Ind1–Ind10 is unknown; do not infer eleven distinct donors. Shared donor does not establish a shared physical stool across all designs.'),
    ('PREDICT versions', 'August and later metadata are retained side by side. Each analysis uses its producing version. Later identifiers resolve three August blanks; they are not retroactively substituted into earlier patient summaries. Recorded visit variables are not interchangeable.'),
    ('CapScan layers', 'DNA and protein/metabolite specimens can differ within a participant-region. The native analysis joins layers at participant-region, not by assuming identical specimens. The 301-run upstream DNA processing roster also includes saliva.'),
    ('Exclusions', 'Outside a saved roster is not evidence of assay failure or clinical dropout. Source-documented causes and reconstructed nonmembership are distinguished. Counts are conditional on the declared source universe; CAMPI is the retained 15-acquisition benchmark, not every acquisition in the original publication.'),
    ('Denominators', 'Roster size is not every feature-specific statistical n. Finite values, annotation and other endpoint-specific criteria still apply. CapScan membership contains separate entity_type rows; count the relevant entity type rather than all rows.'),
    ('Provenance', 'Sources contains exact inspected file locations and fresh whole-file SHA-256 hashes. Original source fields remain unchanged; analysis labels and joins are explicitly described. TSV copies, selection code and validation accompany this workbook.'),
    ('Status', 'Review checkpoint dated 24 September 2026. Frozen source analyses and the integration manuscript are unchanged. This package has not been deposited or published; final supplementary numbering and manuscript fusion remain separate tasks.'),
    ('Scope', 'Roster and source-link checks, not a new raw-spectra search, LFQ analysis or statistical rerun. This is a sample index, not the complete figure source-data package.'),
]
sheet('Read me', [dict(topic=k, explanation=v) for k,v in notes], 'How to interpret this workbook')
sheet('Cohorts', [
    dict(cohort='CAMPI', acquisitions=15, specimens='Defined SIHUMIx mixture; no individual specimen IDs in retained metadata', participants='Not applicable', design='Twelve published workflow labels (S01–S12); S11 contributes four fractions', roster_sheet='CAMPI acquisitions'),
    dict(cohort='Healthy stool', acquisitions=58, specimens='Original specimen IDs unavailable; within-design relationships documented', participants='Ten between-donor labels plus a shared experimental donor of unknown overlap', design='12 reinjections; 15 preparations; 10 spatial sections; 11 longitudinal samples; 10 between-donor samples', roster_sheet='HS acquisitions'),
    dict(cohort='PREDICT', acquisitions=306, specimens='300 explicit tube IDs in August metadata; four acquisitions lack tube links', participants='93 recorded August patients; four acquisitions lack participant links', design='Analysis-specific subsets and later metadata are indexed separately', roster_sheet='PREDICT metadata'),
    dict(cohort='CapScan', acquisitions=331, specimens=292, participants=15, design='Fourteen participants contribute capsule and stool; one contributes stool only', roster_sheet='CAPSCAN acquisitions'),
], 'Four study acquisition rosters; biological denominators are not summed across designs')

indices=[]
for r in read('healthy_campi/analysis_index.tsv'):
    aid=r['analysis_id']; cohort=r['cohort']
    n_people = {'HSTOOL_TechnicalReplicates':1,'HSTOOL_WorkflowReplicates':1,'HSTOOL_OnePiece':1,'HSTOOL_Intra':1,'HSTOOL_Interindividual':10,'HSTOOL_other_eleven_pool':1}.get(aid, 'Not applicable' if cohort=='CAMPI' else 'Unknown overlap between shared donor and Ind1–Ind10')
    indices.append(dict(analysis_id=aid, cohort=cohort, purpose=r['question'], analysis_unit=r['analysis_unit'], n_MS_acquisitions=r['included_acquisitions'], n_recorded_proteomic_specimens='Not supplied as original specimen IDs', n_participants=n_people, n_DNA_runs='Not applicable', n_DNA_specimens='Not applicable', n_participant_region_cells='Not applicable', roster_sheet='HS CAMPI membership', roster_filter='analysis_id = '+aid+'; included = yes', exclusion_sheet='HS excluded' if aid=='HSTOOL_primary' else 'HS CAMPI membership', exclusion_filter='All rows' if aid=='HSTOOL_primary' else 'analysis_id = '+aid+'; included = no', inclusion_rule=r['inclusion_rule'], aggregation='See analysis_unit and source-specific Methods', limitations=r['limits'], provenance=r['source_ids']))
for r in read('predict/ANALYSIS_INDEX.tsv'):
    aid=r['analysis_id']
    indices.append(dict(analysis_id=aid, cohort='PREDICT', purpose=r['human_purpose'], analysis_unit=r['analysis_unit'], n_MS_acquisitions=r['n_acquisitions'], n_recorded_proteomic_specimens=r['n_documented_tubes'], n_participants=r['n_participants'], n_DNA_runs='Not indexed as run-level data in this protein specimen roster', n_DNA_specimens='See complete-stack eligibility for native analysis', n_participant_region_cells='Not applicable', roster_sheet='PREDICT membership', roster_filter='analysis_id = '+aid+'; status = included', exclusion_sheet='PREDICT membership', exclusion_filter='analysis_id = '+aid+'; status = excluded', inclusion_rule=r['inclusion_criterion'], aggregation=r['aggregation'], limitations=r['limitations'], provenance=r['source_ids']))
for r in read('capscan/ANALYSIS_INDEX.tsv'):
    aid=r['analysis_id']
    indices.append(dict(analysis_id=aid, cohort='CapScan', purpose=r['purpose'], analysis_unit=r['biological_unit'], n_MS_acquisitions=r['acquisitions'], n_recorded_proteomic_specimens=r['proteomic_specimens'], n_participants=r['participants'], n_DNA_runs=r['DNA_runs'], n_DNA_specimens=r.get('DNA_specimens', r.get('dna_specimens', 'See DNA entity rows')), n_participant_region_cells=r['participant_region_cells'], roster_sheet='CAPSCAN membership', roster_filter='analysis_id = '+aid+'; select entity_type for measurement/specimen/participant or region', exclusion_sheet='CAPSCAN exclusions', exclusion_filter='analysis_id = '+aid, inclusion_rule=r['inclusion_rule'], aggregation=r['biological_unit'], limitations=r['limitations'], provenance=r['source_file']))
indices.append(dict(analysis_id='CAPSCAN_HISTIDINE_ROUTES', cohort='CapScan', purpose='Histidine terminal routes and internal-boundary inference', analysis_unit='Acquisition input; inference uses endpoint-specific participant or paired-participant profiles', n_MS_acquisitions=331, n_recorded_proteomic_specimens=292, n_participants=15, n_DNA_runs='Not applicable', n_DNA_specimens='Not applicable', n_participant_region_cells=71, roster_sheet='CAPSCAN H acquisitions', roster_filter='All rows; input universe only', exclusion_sheet='CAPSCAN H participants', exclusion_filter='Observed finite H-source and reciprocal endpoint participant sets; other endpoint membership remains in producing sources (M22)', inclusion_rule='Exact A20/A27/A28 revision-4 search acquisition universe; do not apply quantitative330 exclusions', aggregation='Acquisition to specimen to participant-region; paired spatial potential universe has fourteen participants', limitations='The zero-LFQ acquisition still contains terminal evidence. Stool source-region summaries include fifteen participants; finite test n varies. Saved n ranges are not a replay of all 44,010 tests.', provenance='capscan/HISTIDINE_SOURCE_MANIFEST.tsv; capscan/HISTIDINE_ROSTER_SCOPE.md'))
assert len({r['analysis_id'] for r in indices})==len(indices)
sheet('Analyses',indices,'One row per analysis or experimental-design subset; follow roster_sheet and roster_filter')
write_tsv(D/'ANALYSIS_INDEX.tsv',indices)

tables=[
('HS acquisitions','healthy_campi/HSTOOL_acquisitions.tsv','58 healthy-stool acquisitions with documented relationships and explicit unknown identities'),
('HS relationships','healthy_campi/HSTOOL_design_relationships.tsv','Author confirmation and documentary links between designs'),
('HS excluded','healthy_campi/HSTOOL_excluded39.tsv','39 historical differential-centrifugation acquisitions excluded from final study scope'),
('HS CAMPI membership','healthy_campi/analysis_membership.tsv','Included and excluded acquisition membership for healthy-stool/CAMPI analyses'),
('HS transfer','healthy_campi/HSTOOL_transfer_links.tsv','Exact target-to-database-source links; biological independence is separately classified'),
('HS pools','healthy_campi/HSTOOL_pool_links.tsv','Each target and every database-source acquisition in own/all-twelve/other-eleven pools'),
('CAMPI acquisitions','healthy_campi/CAMPI_acquisitions.tsv','Retained deposited acquisitions linked to published protocol and fraction labels'),
('CAMPI search arms','healthy_campi/CAMPI_search_arm_membership.tsv','75 primary and 45 second-engine acquisition/database searches'),
('PREDICT metadata','predict/ACQUISITION_METADATA_LINKS.tsv','August and later explicit metadata links side by side'),
('PREDICT membership','predict/ANALYSIS_MEMBERSHIP.tsv','All 306 acquisitions classified for each PREDICT analysis'),
('PREDICT patients','predict/PATIENT_MEMBERSHIP.tsv','Patient eligibility relative to each documented source population'),
('PREDICT mapping changes','predict/METADATA_MAPPING_CHANGES_AND_UNRESOLVED.tsv','Three later-resolved mappings and the remaining unmapped analytical sample'),
('PREDICT repeated tubes','predict/REPEATED_ACQUISITIONS_OF_RECORDED_TUBES.tsv','Repeated measurements linked to the same recorded tube'),
('PREDICT MG exclusions','predict/PUBLIC_MG10_EXCLUDED.tsv','Ten acquisitions lacking the matched comparator in the producing manifest'),
('PREDICT H patients','predict/HISTIDINE_ENDPOINT_PATIENT_ROSTER.tsv','Patient membership for each of twelve primary histidine endpoints'),
('CAPSCAN H acquisitions','capscan/HISTIDINE_ACQUISITION_ROSTER.tsv','331 histidine search acquisitions; quantitative zero is retained'),
('CAPSCAN H finite n','capscan/HISTIDINE_SAVED_FINITE_N_SUMMARY.tsv','Saved family-specific finite n ranges; no statistical replay'),
('CAPSCAN H participants','capscan/HISTIDINE_FINITE_PARTICIPANT_SETS.tsv','Observed finite participant sets for H-source and reciprocal endpoints only'),
('CAPSCAN acquisitions','capscan/BASE_ACQUISITIONS.tsv','331 acquisition-to-specimen/participant/region links'),
('CAPSCAN membership','capscan/ANALYSIS_MEMBERSHIP.tsv','Included entities; entity_type distinguishes acquisitions, specimens, DNA runs and aggregation units'),
('CAPSCAN exclusions','capscan/EXCLUSIONS.tsv','Documented omissions and source-specific nonmembership'),
('CAPSCAN MS nonmembership','capscan/ACQUISITION_NONMEMBERSHIP.tsv','Every acquisition absent from each MS analysis roster relative to the 331-search design'),
('CAPSCAN DNA materials','capscan/DNA301_SOURCE_CATEGORIES.tsv','Upstream DNA specimen categories: 272 gut/stool and 29 saliva'),
('PREDICT 85 to 77','predict/EDS85_TO_NATIVE77_PATIENT_EXCLUSIONS.tsv','Eight patients without any complete-stack specimen; matrix row absence is not a known laboratory cause'),
('CAPSCAN ID crosswalk','capscan/IDENTIFIER_CROSSWALK.tsv','Exact linked biology/cross-omics specimen namespaces'),
('CAPSCAN MG pairs','capscan/BENCHMARK282_ACQUISITIONS.tsv','Matched benchmark pairs preserving A16 and biology specimen labels'),
('CAPSCAN DNA runs','capscan/DNA301_RUN_SPECIMEN_ROSTER.tsv','301 upstream completed DNA runs, including saliva; native subset is in membership'),
('CAPSCAN DNA eligibility','capscan/DNA_CELL_ELIGIBILITY.tsv','Participant-region DNA availability across functional representations'),
('CAPSCAN regional pairs','capscan/REGIONAL_PAIRED_LINKS.tsv','Four capsule-versus-stool pairs for each of fourteen participants'),
]
# Optional precise full-stack availability table supplied by the PREDICT audit.
for fn in ['LATER268_ASSAY_ELIGIBILITY.tsv','FULLSTACK_ELIGIBILITY.tsv','ASSAY_INTERSECTION_ELIGIBILITY.tsv','NATIVE77_ELIGIBILITY.tsv']:
    if (O/'predict'/fn).exists():
        tables.append(('PREDICT stack eligibility','predict/'+fn,'Recorded assay presence and complete-stack eligibility'))
        break
for name,rel,desc in tables:
    sheet(name,read(rel),desc,rel)

sources=[]
for sub,manifest in [('healthy_campi','healthy_campi/SOURCE_MANIFEST.tsv'),('predict','predict/SOURCE_MANIFEST.tsv'),('capscan','capscan/SOURCE_MANIFEST.tsv'),('capscan_histidine','capscan/HISTIDINE_SOURCE_MANIFEST.tsv')]:
    for r in read(manifest):
        p=Path(r.get('path',r.get('sourcefile',r.get('source_file',''))))
        actual=digest(p)
        assert actual==r['sha256'],str(p)
        sources.append(dict(workstream=sub,source_id=r.get('source_id',''),source_path=str(p),sha256=actual,bytes=p.stat().st_size,local_snapshot=r.get('snapshot','')))
sheet('Sources',sources,'Exact inspected inputs; all recorded hashes freshly verified during workbook build')
write_tsv(D/'SOURCE_MANIFEST.tsv',sources)
author=[]
for line in (O/'AUTHOR_CONFIRMATIONS.jsonl').read_text().splitlines():
    r=json.loads(line)
    author.append({k:'; '.join(v) if isinstance(v,list) else v for k,v in r.items() if k != 'verbatim'})
sheet('Author confirmations',author,'Author-confirmed relationships with explicit scope; verbatim record retained in AUTHOR_CONFIRMATIONS.jsonl')
sheet('Sheet directory',sheet_records.copy(),'Location and original table for each workbook sheet')

# Add navigation links without altering stored identifiers or field values.
for name in ['Analyses','Cohorts','Sheet directory']:
    ws=wb[name]; headers={c.value:c.column for c in ws[1]}
    for key in ['roster_sheet','exclusion_sheet','sheet']:
        if key in headers:
            for row in range(2,ws.max_row+1):
                cell=ws.cell(row,headers[key])
                if cell.value in wb.sheetnames:
                    cell.hyperlink="#'%s'!A1" % cell.value
                    cell.font=Font(color='0563C1',underline='single')
for ws in wb:
    ws.sheet_properties.tabColor = '497E76' if ws.title.startswith('HS') else '517CA4' if ws.title.startswith('PREDICT') else '98689A' if ws.title.startswith('CAPSCAN') else 'B58A43'
wb.save(D/'Sample_rosters.xlsx')

# Confirm no Excel type conversion, truncation or lost rows on round trip.
check=load_workbook(D/'Sample_rosters.xlsx',read_only=True,data_only=False)
for name,values in expected.items():
    actual=[[str(v) if v is not None else '' for v in row] for row in check[name].iter_rows(values_only=True)]
    assert actual==values,name
check.close()
counts={k:len(read(v)) for k,v in [('CAMPI','healthy_campi/CAMPI_acquisitions.tsv'),('Healthy stool','healthy_campi/HSTOOL_acquisitions.tsv'),('PREDICT','predict/ACQUISITION_METADATA_LINKS.tsv'),('CapScan','capscan/BASE_ACQUISITIONS.tsv')]}
assert counts=={'CAMPI':15,'Healthy stool':58,'PREDICT':306,'CapScan':331}
assert sum(counts.values())==710
for r in indices:
    aid=r['analysis_id']; n=int(r['n_MS_acquisitions'])
    if r['cohort'] in ['HSTOOL','CAMPI']:
        found=[x for x in read('healthy_campi/analysis_membership.tsv') if x['analysis_id']==aid and x['included']=='yes']
    elif r['cohort']=='PREDICT':
        found=[x for x in read('predict/ANALYSIS_MEMBERSHIP.tsv') if x['analysis_id']==aid and x['status']=='included']
    elif aid=='CAPSCAN_HISTIDINE_ROUTES':
        found=read('capscan/HISTIDINE_ACQUISITION_ROSTER.tsv')
        base={x['acquisition_id']:x for x in read('capscan/BASE_ACQUISITIONS.tsv')}
        assert all(x['participant_id']==base[x['acquisition_id']]['participant_id'] and x['histidine_specimen_id']==base[x['acquisition_id']]['specimen_id_biology'] and x['region']==base[x['acquisition_id']]['region'] for x in found)
        assert sum(x['in_quantitative330_roster']=='False' for x in found)==1
    else:
        found=[x for x in read('capscan/ANALYSIS_MEMBERSHIP.tsv') if x['analysis_id']==aid and x['entity_type']=='acquisition']
    assert len(found)==n,(aid,len(found),n)
    assert len({x['acquisition_id'] for x in found})==n,aid
report=dict(status='PASS',cohort_acquisitions=counts,total_acquisitions=710,analysis_index_rows=len(indices),workbook_sheets=len(wb.sheetnames),workbook_exact_text_roundtrip=True,analysis_acquisition_counts_checked=True,source_files_rehashed=len(sources),workbook_sha256=digest(D/'Sample_rosters.xlsx'),scope='Roster/package checks; no new scientific analysis. Original metadata identities and versions retained.')
(O/'WORKBOOK_VALIDATION.json').write_text(json.dumps(report,indent=2)+'\n')
# Include source-derived TSVs and builders, not a duplicate of large underlying assay matrices.
with zipfile.ZipFile(D/'Sample_rosters_tables_and_provenance.zip','w',zipfile.ZIP_DEFLATED) as z:
    for sub in ['healthy_campi','predict','capscan']:
        for p in sorted((O/sub).iterdir()):
            if p.is_file() and p.suffix in ['.tsv','.json','.py','.md']:
                z.write(p,str(p.relative_to(O)))
    for p in [O/'AUTHOR_CONFIRMATIONS.jsonl',O/'WORKBOOK_VALIDATION.json',Path(__file__)]:
        z.write(p,str(p.relative_to(O)))
    for p in [D/'ANALYSIS_INDEX.tsv',D/'SOURCE_MANIFEST.tsv']:
        z.write(p,p.name)
print(json.dumps(report,indent=2))
