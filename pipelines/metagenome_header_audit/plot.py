"""Plot cumulative distinct database sequences per added specimen (fixed order and order-averaged) with per-specimen peptide IDs, and the header audit, from full/<cohort>/ outputs of worker.py. Checks worker.py sha256 against COMPLETE.json. Writes figures/ and figures/SUMMARY.json."""
from pathlib import Path
import json,csv,hashlib
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.ticker import FuncFormatter
from docx import Document
from docx.shared import Inches
H=Path(__file__).resolve().parent;O=H/'figures';O.mkdir(exist_ok=True)
colors={'FL_public':'#D55E00','MG':'#0072B2'};names={'FL_public':'FastaLake (public resources)','MG':'Matched metagenome + human'}
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none','pdf.fonttype':42})
def rows(p):return list(csv.DictReader(p.open(),delimiter='\t'))
receipts={};data={};headers={}
for c in ['PREDICT','CAPSCAN']:
 d=H/'full'/c;r=json.loads((d/'COMPLETE.json').read_text());assert r['status']=='PASS' and r['mode']=='full';assert r['code_sha256']==hashlib.sha256((H/'worker.py').read_bytes()).hexdigest();receipts[c]=r;data[c]=rows(d/'accumulation.tsv');headers[c]=rows(d/'header_audit.tsv')
summary=[];artifacts=[]
def save(fig,name):
 for ext in ['pdf','svg','png']:
  p=O/(name+'.'+ext);fig.savefig(p,dpi=220,bbox_inches='tight');artifacts.append(p)
 plt.close(fig)
for scale in ['linear','log']:
 fig,axs=plt.subplots(1,2,figsize=(13,4.5))
 for ax,c in zip(axs,['PREDICT','CAPSCAN']):
  right=ax.twinx();right.spines['right'].set_visible(True)
  for arm in colors:
   r=[r for r in data[c] if r['arm']==arm];k=np.array([int(x['k']) for x in r]);db=np.array([int(x['cumulative_sequences_fixed_order']) for x in r]);ids=np.array([int(x['peptide_ids_in_this_specimen']) for x in r]);ax.plot(k,db,color=colors[arm],lw=2);right.plot(k,ids,color=colors[arm],lw=.55,alpha=.6,ls=':',marker='o',ms=2.3)
   if scale=='linear':summary.append({'cohort':c,'arm':arm,'specimens':len(k),'acquisitions':receipts[c]['acquisitions'],'participants':receipts[c]['participants'],'union_sequences':int(db[-1]),'median_peptide_ids_per_specimen':float(np.median(ids))})
  ax.set_yscale(scale);ax.set_xlabel('Specimens added (fixed specimen-ID order)');ax.set_ylabel('Cumulative distinct database sequences');right.set_ylabel('Peptide IDs in the individual specimen');ax.yaxis.set_major_formatter(FuncFormatter(lambda y,pos:f'{y/1e6:g}M' if y>=1e6 else f'{y/1e3:g}k'));right.yaxis.set_major_formatter(FuncFormatter(lambda y,pos:f'{y/1e3:g}k'));right.set_ylim(bottom=0)
  rr=receipts[c];ax.set_title(f"{c.title() if c=='CAPSCAN' else c}\n{rr['specimens']} specimens · {rr['participants']} participants");ax.grid(axis='y',alpha=.2)
 fig.legend(handles=[Line2D([0],[0],color=colors[a],lw=2,label=names[a]) for a in colors]+[Line2D([0],[0],color='black',lw=2,label='Database accumulation (left axis)'),Line2D([0],[0],color='black',ls=':',marker='o',ms=3,label='Individual-specimen peptide IDs (right axis)')],loc='lower center',ncol=2,bbox_to_anchor=(.5,-.15),frameon=False)
 fig.tight_layout();save(fig,'sequence_accumulation_dual_axis_'+scale)
fig,axs=plt.subplots(2,2,figsize=(11,7))
for col,c in enumerate(['PREDICT','CAPSCAN']):
 for arm in colors:
  r=[r for r in data[c] if r['arm']==arm];k=np.array([int(x['k']) for x in r]);mean=np.array([float(x['mean_cumulative_sequences_random_order']) for x in r]);axs[0,col].plot(k,mean,color=colors[arm],lw=2,label=names[arm]);axs[1,col].plot(k,np.diff(np.r_[0,mean]),color=colors[arm],lw=2)
 axs[0,col].set_title(c.title() if c=='CAPSCAN' else c)
 for row in range(2):axs[row,col].set_yscale('log');axs[row,col].set_xlabel('Number of specimens');axs[row,col].grid(axis='y',alpha=.2)
 axs[0,col].set_ylabel('Expected cumulative distinct sequences');axs[1,col].set_ylabel('Expected new sequences per added specimen')
axs[0,0].legend(frameon=False);fig.tight_layout();save(fig,'sequence_accumulation_order_averaged')
fig,axs=plt.subplots(1,2,figsize=(12,4.6));labels=['FastaLake\nas searched','Metagenome\nas searched','Metagenome\noriginal full headers','Metagenome\nassembly-local names'];keys=['FL_public_as_searched','MG_as_searched','MG_original_prefixed_headers','MG_assembly_local_headers']
for ax,c in zip(axs,['PREDICT','CAPSCAN']):
 lookup={r['namespace']:r for r in headers[c]};x=np.arange(4)
 for j,(metric,label,color) in enumerate([('alias_rate_among_shared_sequences','Shared sequence has multiple headers','#009E73'),('collision_rate_among_shared_headers','Shared header has multiple sequences','#CC79A7')]):
  vals=[float(lookup[k][metric])*100 if lookup[k][metric] else np.nan for k in keys];bars=ax.bar(x+(j-.5)*.36,vals,width=.36,color=color,label=label)
  for b,v in zip(bars,vals):
   if np.isfinite(v):ax.text(b.get_x()+b.get_width()/2,v+1,f'{v:.1f}%',ha='center',fontsize=8)
 ax.set_xticks(x,labels);ax.set_ylim(0,112);ax.set_ylabel('Inconsistency among shared identities (%)');ax.set_title(c.title() if c=='CAPSCAN' else c);ax.grid(axis='y',alpha=.15)
axs[0].legend(frameon=False,fontsize=8,loc='upper left',bbox_to_anchor=(0,-.25));fig.tight_layout();save(fig,'header_inconsistency_audit')
# Exact observed cross-study union; no extrapolated saturation parameter.
for arm in colors:
 a=np.load(H/'full/PREDICT'/(arm+'_sequence_incidence.npz'))['keys'];b=np.load(H/'full/CAPSCAN'/(arm+'_sequence_incidence.npz'))['keys'];shared=len(np.intersect1d(a,b,assume_unique=True));summary.append({'cohort':'cross_study','arm':arm,'predict_sequences':len(a),'capscan_sequences':len(b),'shared_sequences':shared,'combined_union_sequences':len(a)+len(b)-shared});del a,b
caption='Paired public-resource FastaLake and matched-metagenome-plus-human databases from the current manuscript: 296 PREDICT and 282 CapScan acquisitions. Technical/aliquot acquisitions are combined within frozen specimen IDs; participant and specimen counts are shown. Protein identity is the exact amino-acid sequence represented by full SHA256; I/L is not collapsed for database sequences. Solid lines show cumulative distinct sequences in a declared fixed specimen-ID order (left axis); dotted points show actual canonical peptide IDs in each individual specimen, not cumulative identifications (right axis). Peptide IDs are the union of rank-one target identifications at peptide q≤0.01 and lengths5–50, with modifications stripped and I/L treated as equivalent, matching the primary manuscript definition. The order-averaged companion is the exact expected sequence union under uniformly random specimen order, computed from specimen-incidence frequencies; it does not supply a confidence interval or infer a finite biological upper bound. Header audit denominators are identities shared across at least two specimens; full counts and all-identity denominators are in the source tables. Original full and assembly-local metagenome headers are kept distinct from the actual hashed search headers. These are database-sequence and peptide-identification counts, not search-engine candidate counts or calibrated whole-workflow FDR.'
(O/'CAPTIONS.md').write_text(caption+'\n');(O/'SUMMARY.json').write_text(json.dumps(summary,indent=2)+'\n')
doc=Document();doc.add_heading('FastaLake — empirical database accumulation and header audit',0);doc.add_paragraph('Figure proposals for Peter’s review. Candidate-per-spectrum endpoint is separate and remains pending. Figures are not yet inserted into the manuscript.');doc.add_paragraph(caption)
for name in ['sequence_accumulation_dual_axis_linear','sequence_accumulation_dual_axis_log','sequence_accumulation_order_averaged','header_inconsistency_audit']:
 doc.add_heading(name.replace('_',' '),1);doc.add_picture(str(O/(name+'.png')),width=Inches(6.4))
doc.save(O/'Sequence_accumulation_figure_proposals.docx');artifacts.append(O/'Sequence_accumulation_figure_proposals.docx')
(O/'COMPLETE.json').write_text(json.dumps({'status':'FIGURES_GENERATED_PENDING_VISUAL_REVIEW','cohorts':receipts,'source_summary':summary,'candidate_per_spectrum':'pending definition and measurement','files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in artifacts}},indent=2)+'\n')
print('FIGURES GENERATED',O,flush=True)
