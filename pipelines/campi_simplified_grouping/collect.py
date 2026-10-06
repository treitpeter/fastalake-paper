#!/usr/bin/env python3
"""Validate all scheduled tasks and collect their reaggregated differences."""

import argparse
import csv
import gzip
import hashlib
import json
import math
from pathlib import Path

from openpyxl import Workbook

LEVELS = (
    "OG",
    "KEGG_ko",
    "KEGG_Pathway",
    "KEGG_Module",
    "KEGG_Reaction",
    "COG_category",
    "CAZy",
    "Pfam",
    "GOs",
    "BRITE",
    "EC",
    "taxon",
)


def validate_task(root, task):
    folder = root / "cohorts" / task["id"]
    complete = folder / "COMPLETE.json"
    if not complete.is_file():
        raise ValueError("COMPLETE.json is missing; inspect the task log and FAILED.json")
    summary = json.loads(complete.read_text())
    names = Path(task["samples_file"]).read_text().splitlines()
    if summary["status"] != "PASS" or set(summary["acquisition_names"]) != set(names):
        raise ValueError("Completed population does not match the scheduled roster")
    if summary["acquisitions"] != len(names) or len(names) != task["acquisitions"]:
        raise ValueError("Completed acquisition count does not match the schedule")
    if not all(summary["checks"].values()):
        raise ValueError("An accounting/compatibility check failed")
    expected = [
        folder / "direct" / f
        for f in [
            "study_group_matrix.tsv",
            "study_group_long.tsv",
            "feature_assignments.tsv.gz",
            "input_manifest.json",
            "summary.json",
        ]
    ]
    expected += [
        folder / "reaggregated" / f
        for f in [
            "REPORT.md",
            "SUMMARY.json",
            "original_study_group_matrix.tsv.gz",
            "simplified_study_group_matrix.tsv.gz",
            "median_diff.tsv",
            "per_acquisition.tsv",
            "group_diff.tsv.gz",
            "PROVENANCE.json",
            "fixed_pipeline_metrics.tsv",
            "KO_leader_diff.tsv.gz",
            "KO_taxon_long.tsv.gz",
            "annotation_resolution.tsv",
            "representative_annotations.tsv.gz",
            "functional_profile_agreement.tsv",
            "group_detection.tsv.gz",
        ]
    ]
    for level in LEVELS:
        for mode in ("full", "split"):
            expected += [
                folder / "reaggregated/functional" / f"{m}_{level}_{mode}_matrix.tsv.gz"
                for m in ("original", "simplified")
            ]
            expected.append(folder / "reaggregated/functional" / f"{level}_{mode}_diff.tsv.gz")
    for path in expected:
        if not path.is_file() or not path.stat().st_size:
            raise ValueError(f"Required reaggregated output missing/empty: {path}")
    for method in ("original", "simplified"):
        matrix = folder / "reaggregated" / f"{method}_study_group_matrix.tsv.gz"
        with gzip.open(matrix, "rt") as f:
            header = next(csv.reader(f, delimiter="\t"))
            if header[1:] != sorted(names):
                raise ValueError("Group matrix acquisition columns do not match the cohort")
            for row in csv.reader(f, delimiter="\t"):
                if len(row) != len(header):
                    raise ValueError("Malformed group matrix row")
    with (folder / "reaggregated/per_acquisition.tsv").open() as f:
        totals = {
            (r["method"], r["sample"]): float(r["intensity"])
            for r in csv.DictReader(f, delimiter="\t")
        }
    for method in ("original", "simplified"):
        for level in LEVELS:
            for mode in ("full", "split"):
                path = folder / "reaggregated/functional" / f"{method}_{level}_{mode}_matrix.tsv.gz"
                with gzip.open(path, "rt") as f:
                    reader = csv.reader(f, delimiter="\t")
                    header = next(reader)
                    if header[1:] != sorted(names):
                        raise ValueError(f"Functional matrix cohort mismatch: {path}")
                    sums = [0.0] * len(names)
                    for row in reader:
                        if len(row) != len(header):
                            raise ValueError(f"Malformed functional matrix: {path}")
                        for i, value in enumerate(row[1:]):
                            if value:
                                value = float(value)
                                if not math.isfinite(value) or value <= 0:
                                    raise ValueError(f"Invalid functional intensity: {path}")
                                sums[i] += value
                    if mode == "split" and any(
                        not math.isclose(v, totals[(method, s)], rel_tol=1e-9, abs_tol=1e-6)
                        for s, v in zip(sorted(names), sums)
                    ):
                        raise ValueError(f"Emitted split matrix fails conservation: {path}")
    return summary, len(expected)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()
    manifest = json.loads(args.manifest.read_text())
    root = Path(manifest["output_root"])
    output = root / "collection"
    output.mkdir(exist_ok=False)
    statuses, diffs = [], []
    workbook = Workbook(write_only=True)
    for task in manifest["tasks"]:
        try:
            summary, count = validate_task(root, task)
            statuses.append(
                dict(
                    task=task["id"],
                    cohort=task["cohort"],
                    arm=task["arm"],
                    status="PASS",
                    acquisitions=summary["acquisitions"],
                    checked_outputs=count,
                    error="",
                )
            )
            folder = root / "cohorts" / task["id"] / "reaggregated"
            records = list(csv.DictReader((folder / "median_diff.tsv").open(), delimiter="\t"))
            diffs.extend(
                dict(task=task["id"], cohort=task["cohort"], arm=task["arm"], **r) for r in records
            )
            for suffix, filename in [
                ("medians", "median_diff.tsv"),
                ("samples", "per_acquisition.tsv"),
            ]:
                sheet = workbook.create_sheet(f"{task['index']}_{task['cohort']}_{suffix}"[:31])
                sheet.freeze_panes = "A2"
                with (folder / filename).open() as f:
                    for number, row in enumerate(csv.reader(f, delimiter="\t")):
                        label_columns = 1 if suffix == "medians" else 2
                        sheet.append(
                            row
                            if number == 0
                            else [
                                v if i < label_columns else float(v) if v else None
                                for i, v in enumerate(row)
                            ]
                        )
        except Exception as error:
            statuses.append(
                dict(
                    task=task["id"],
                    cohort=task["cohort"],
                    arm=task["arm"],
                    status="FAIL",
                    acquisitions=task["acquisitions"],
                    checked_outputs=0,
                    error=str(error),
                )
            )
    for filename, records in [("task_status.tsv", statuses), ("all_median_differences.tsv", diffs)]:
        if records:
            with (output / filename).open("w", newline="") as f:
                w = csv.DictWriter(f, fieldnames=list(records[0]), delimiter="\t")
                w.writeheader()
                w.writerows(records)
    ws = workbook.create_sheet("Task status")
    ws.append(list(statuses[0]))
    for row in statuses:
        ws.append(list(row.values()))
    workbook.save(output / "ALL_COHORT_DIFFS.xlsx")
    passed = all(r["status"] == "PASS" for r in statuses)
    lines = [
        "# Scheduled manuscript regrouping",
        "",
        f"Overall status: **{'PASS' if passed else 'FAIL — incomplete tasks remain'}**.",
        "",
        "| Task | Cohort | Arm | Acquisitions | Status |",
        "|---|---|---|---:|---|",
    ]
    for row in statuses:
        link = f"../cohorts/{row['task']}/reaggregated/REPORT.md"
        lines.append(
            f"| [{row['task']}]({link}) | {row['cohort']} | {row['arm']} | "
            f"{row['acquisitions']} | {row['status']} |"
        )
    lines += [
        "",
        "[Combined difference workbook](ALL_COHORT_DIFFS.xlsx) · "
        "[All median differences](all_median_differences.tsv) · [Task status](task_status.tsv)",
        "",
        "Each successful task includes both original and simplified group aggregates, all eleven",
        "functional annotation levels plus taxon, full/split rollups, and explicit differences.",
        "Input searches are retained; this is a post-search workflow.",
        "",
        "Tomorrow: review every affected figure's source tables and captions. Group counts,",
        "peptide support, coverage, quantification, functional contributors and biological",
        "conclusions may need to be rebuilt. Identification-only quantities retain their units.",
        "Artwork and manuscripts have not been overwritten. Species labels and group FDR remain",
        "separate questions from sequence compatibility and intensity accounting.",
        "",
    ]
    lines.extend(f"- FAILED {r['task']}: {r['error']}" for r in statuses if r["status"] != "PASS")
    (output / "REPORT.md").write_text("\n".join(lines))
    result = dict(
        status="PASS" if passed else "FAIL",
        tasks=statuses,
        manifest_sha256=hashlib.sha256(args.manifest.read_bytes()).hexdigest(),
        code_commit=manifest["code_commit"],
        expected_tasks=len(manifest["tasks"]),
    )
    (output / "VALIDATION.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2), flush=True)
    if not passed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
