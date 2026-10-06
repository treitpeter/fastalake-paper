from pathlib import Path
from datetime import datetime,timezone
import hashlib,shutil,subprocess
import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from analyze import HERE as H,load,dump,verify

J=H.parents[1];stamp=datetime.now(timezone.utc).isoformat()
r=load(H/'mechanism_v2/COMPLETE.json');assert r['acquisitions']==15
for x in r['outputs']:verify(x)
for name in ['gain_availability.tsv','gain_spectrum_fates.tsv','new_quantified_peptides.tsv','destination_group_support.tsv','upstream_prediction_stages.tsv']:
 a=pd.read_csv(H/'mechanism'/name,sep='\t').replace('digest_compatible_not_accepted_FL','digest_compatible_in_FL_fasta')
 b=pd.read_csv(H/'mechanism_v2'/name,sep='\t')
 # Renaming a grouping key changes lexical row order, but not the results.
 pd.testing.assert_frame_equal(a.sort_values(list(a)).reset_index(drop=True),b.sort_values(list(b)).reset_index(drop=True))
primary=pd.read_csv(H/'summary/primary_counts.tsv',sep='\t').query('min_length==5').pivot(index='sample',columns='arm',values='peptides')
a=pd.read_csv(H/'mechanism_v2/gain_availability.tsv',sep='\t')
q=pd.read_csv(H/'mechanism_v2/new_quantified_peptides.tsv',sep='\t')
f=pd.read_csv(H/'mechanism_v2/gain_spectrum_fates.tsv',sep='\t')
g=pd.read_csv(H/'mechanism_v2/destination_group_support.tsv',sep='\t')
for (sample,dest),x in a.groupby(['sample','destination']):
 assert x.gained_peptides.sum()==x.total_gained.iloc[0]
 assert x.total_gained.nunique()==x.total_lost.nunique()==1
 assert x.total_gained.iloc[0]-x.total_lost.iloc[0]==primary.loc[sample,dest]-primary.loc[sample,'FL']
for _,x in q.groupby(['sample','destination']):assert x.peptides.sum()==x.total_destination_only_quantified.iloc[0]
for _,x in f.groupby(['sample','destination']):assert x.spectra.sum()==x.total_new_peptide_spectra.iloc[0]
for _,x in g.groupby(['sample','destination']):assert x.groups.sum()==x.total_groups.iloc[0]
stats={}
for dest in ['FLR','FLRM']:
 x=a[a.destination==dest];t=x.drop_duplicates('sample');z=q[q.destination==dest]
 stats[dest]=dict(gained=int(x.gained_peptides.sum()),lost=int(t.total_lost.sum()),net=int(t.net_change.sum()),unavailable=int(x[x.availability!='digest_compatible_in_FL_fasta'].gained_peptides.sum()),absent_sequence=int(x[x.availability=='sequence_absent_from_FL_fasta'].gained_peptides.sum()),searchable=int(x[x.availability=='digest_compatible_in_FL_fasta'].gained_peptides.sum()),new_quantified=int(z.peptides.sum()),new_identified_and_quantified=int(z[z.category=='newly_identified_and_quantified'].peptides.sum()),previously_identified_new_quantified=int(z[z.category=='already_identified_FL_newly_quantified'].peptides.sum()))
b=stats['FLR'];pct=100*b['unavailable']/b['gained'];pq=100*b['new_identified_and_quantified']/b['new_quantified']
up=pd.read_csv(H/'mechanism_v2/upstream_prediction_stages.tsv',sep='\t').groupby('stage').sum(numeric_only=True)
u=up.loc['absent_from_acquisition_candidates'];assert int(u.peptides)==31843 and int(u.direct_prediction_lookup)==0
mechanism=(
 f'We traced every additional accepted peptide and quantified feature to the selected FASTA, the spectrum-level result and the protein-group assignment in the current 15 SIHUMIx acquisitions. REF add-back yielded {b["gained"]:,} gained and {b["lost"]:,} lost peptide-acquisition observations (net {b["net"]:,}); these are observations summed across acquisitions, not pooled unique peptides. '
 f'Of the gains, {b["absent_sequence"]:,} (95.4%) were absent as sequences from the original FL FASTA and 1,616 (3.1%) occurred only in digestion-incompatible contexts. Thus {pct:.1f}% were not searchable under the original digestion rules: adding proteins restores sequence candidates that the original search could not report. Only {b["searchable"]:,} gains (1.6%) were already searchable. This identifies sequence availability as the dominant mechanism of the add-back gains, without declaring every gained match correct. '
 'For REF-supported observations missing from FL, the upstream audit located 31,843 before candidate retrieval: 19,241 had no exact de novo prediction, 2,679 had exact predictions only among rows failing length/positive-score/complexity criteria, and 9,923 had qualifying predictions but none retained in the top-scoring fraction. Proteins can also be retrieved by other peptide queries, so these direct-query categories describe the observed losses rather than a universal one-query/one-protein causal rule. A further 17,516 REF peptide observations had digestion-compatible candidate carriers that were discarded during compact-database selection. These stages are distinct from the 1,531 observations already searchable in the selected database but not accepted. '
 'The spectrum audit retains adverse effects: added candidates also changed winners and caused confidence-threshold losses; a larger database did not improve every assignment. '
 f'Downstream, {b["new_identified_and_quantified"]:,}/{b["new_quantified"]:,} ({pq:.1f}%) newly quantified peptide observations were also newly identified; only {b["previously_identified_new_quantified"]} had already been identified in FL. Extra quantitative evidence therefore enters overwhelmingly through newly accepted identifications. Groups supported by the shared quantified peptide set actually decreased after add-back in all 15 acquisitions. The net group increase comes from arm-specific quantified evidence, rather than simply splitting shared evidence into more groups. The second add-back arm gives the same qualitative result. '
 'Why the acquisitions differ is a separate, connected question: their spectral sampling and sequence-selection outputs differ. The current per-acquisition pipeline table traces spectra through qualifying queries, candidates, selected entries, accepted peptides and quantified groups. Distinct-query counts do not rise monotonically with spectrum count, and gradient and load co-vary in the Kleiner series. Equal-MS2 post-search standardization reduces the FL peptide range from 3.38-fold to 1.98-fold; the fixed-load CEA 60-to-120-minute pair changes from a 37.6% observed peptide gain to a 0.7% lower expected yield at equal MS2 depth. Those conditional results support a sampling-depth contribution but do not refit the searches or establish a replicated causal gradient effect. Individual S11 fractions also differ from whole-sample runs: their exact union has 42,229 FL peptides and 4,585 fixed groups. Across all acquisitions, database effects are real but do not eliminate the spread; the peptide range remains 3.19-fold with REF add-back. We therefore do not attribute the full spread solely to database size, preparation, or sampling depth.'
)
addition=(f'Post-search subsampling to equal MS2 counts narrowed the between-acquisition peptide-yield range from 3.38-fold to 1.98-fold, while reference-proteome add-back increased quantified protein-group yields by a median 36.8%. Of the additional accepted peptide observations after add-back, {pct:.1f}% were not searchable in the original selected database, and newly identified peptides supplied {pq:.1f}% of the newly quantified peptide observations.')
review=load(H/'SCIENTIFIC_REVIEW.json')
review.update(status='MECHANISTIC_FINDINGS_REVIEWED_WORDING_PENDING',mechanism=mechanism,proposed_addition=addition,mechanism_statistics=stats,mechanism_job='6031428',mechanism_label_correction=str(H/'MECHANISM_LABEL_CORRECTION.json'),mechanism_validation=dict(primary_gain_loss_partitions=30,new_quantified_partitions=30,spectrum_partitions=30,group_partitions=30,corrected_run_numerically_identical=True))
dump(H/'MECHANISTIC_REVIEW.json',review)
B=H/'before_mechanistic_response';B.mkdir(exist_ok=False)
paths=[J/'POINT_BY_POINT.json',J/'THREE_AUTHOR_OVERLAY.json',J/'top_down_review_2026-09-30/P019_JESS_LAYER_PROPOSAL.json',H.parents[3]/'FastaLake_v17_paragraph_walkthrough.docx']
for p in paths:shutil.copy2(p,B/p.name)
items=load(paths[0]);item=next(x for x in items if x['id']=='JAG17')
item.setdefault('response_history',[]).append(dict(utc=stamp,previous_response=item['response'],reason='Peter requests why/how, now traced empirically.'))
item.update(response=mechanism+'\n\n'+review['response'],candidate_text=item['current_text']+' '+addition,status='mechanistic_evidence_reviewed_proposed_wording_pending',evidence=str(H/'MECHANISTIC_REVIEW.json')+'; '+str(H/'JAG17_SIHUMIx_empirical_review.docx'),remaining='Review the mechanism-based two-sentence addition with Peter; retain gains/losses and conditional post-search depth interpretation. Community-complexity effects remain untested because this analysis is explicitly SIHUMIx-only.',Peter_approval=False)
dump(paths[0],items)
o=load(paths[1]);next(x for x in o['index'] if x['id']=='JAG17')['assessment']='Mechanism audit passed: sequence availability, spectrum results, identified-to-quantified evidence and groups traced on current SIHUMIx data. '+item['remaining'];dump(paths[1],o)
p=load(paths[2]);p.setdefault('proposal_history',[]).append(dict(status=p['status'],proposed_text=p['proposed_text'],addition=p['addition']))
p.update(status='MECHANISM_BASED_PROPOSAL_PENDING_PETER_REVIEW',proposed_text=item['candidate_text'],addition=addition,mechanism_evidence=str(H/'MECHANISTIC_REVIEW.json'),manuscript_changed=False,utc=stamp);dump(paths[2],p)
packet=H/'JAG17_SIHUMIx_empirical_review.docx';shutil.copy2(packet,B/packet.name)
doc=Document();doc.add_heading('JAG17 — why SIHUMIx yields vary',0);doc.add_paragraph('Current v17 results, 15 acquisitions. Mechanism jobs 6031420 and corrected-label repeat 6031428 completed. No manuscript text changed.')
doc.add_heading('Mechanism traced empirically',1);doc.add_paragraph(mechanism)
doc.add_heading('Proposed manuscript addition — pending Peter',1);doc.add_paragraph(addition)
doc.add_heading('Full acquisition/database analysis',1);doc.add_paragraph(review['response'])
doc.add_heading('Interpretation limits',1)
for line in review['limitations']:doc.add_paragraph(line,style='List Bullet')
doc.add_heading('Pipeline trace',1);pipe=pd.read_csv(H/'summary/pipeline_per_acquisition.tsv',sep='\t');cols=['sample','ms2','distinct_queries','selected_entries','peptides','quantified_groups'];t=doc.add_table(rows=1,cols=len(cols));t.style='Light Shading Accent 1'
for c,k in zip(t.rows[0].cells,cols):c.text=k
for _,row in pipe.iterrows():
 for c,k in zip(t.add_row().cells,cols):c.text=str(row[k])
doc.add_paragraph('Evidence: '+str(H/'MECHANISTIC_REVIEW.json'));doc.save(packet)
doc=Document(paths[3]);anchor=next(p for p in doc.paragraphs if p.text.startswith('P019 ·'))
for text in ['JAG17 — mechanism review; current proposal supersedes earlier count-only draft',mechanism,'Current proposed two-sentence addition — awaiting Peter:',addition,'Current full P019 proposal:',item['candidate_text'],'Current evidence: '+str(H/'MECHANISTIC_REVIEW.json')+'; no manuscript edit.']:
 n=OxmlElement('w:p');anchor._p.addnext(n);anchor=Paragraph(n,anchor._parent);anchor.add_run(text)
doc.save(paths[3])
ms=H.parents[3]/'FastaLake_v17_Peter_review_clean.docx';before=hashlib.sha256(ms.read_bytes()).hexdigest()
subprocess.run(['python3',str(J/'build_living_response.py')],check=True)
assert hashlib.sha256(ms.read_bytes()).hexdigest()==before==review['current_manuscript_sha256']
entry='\n\n## JAG17 why/how traced — '+stamp+'\nMechanism job6031428 passed for all15 current SIHUMIx acquisitions and both add-backs; original6031420 outputs preserved with an availability-label correction (no numeric changes). A14 digest functions reused; all reference final-FASTA parent/digest counts reproduced. REF add-back:52720 gained,18980 lost peptide-acquisition observations (net33740).98.4% of gained observations lacked searchable FL sequence;99.9% of newly quantified observations were newly identified. Query/filter/candidate-selection stages, shared group support and the spectra→queries→candidates→selection→IDs→groups chain traced. Evidence: '+str(H/'MECHANISTIC_REVIEW.json')+'. Word packet, Jess response and walkthrough updated with reasoning, limitations and proposed two sentences. Manuscript unchanged; next endpoint Peter wording review. Community complexity untested per SIHUMIx-only scope. Prior KC/Misha/Peter decisions preserved.\n'
for p in [J/'THREE_AUTHOR_REVIEW_HANDOFF.md',H/'STATUS.md']:
 with p.open('a') as f:f.write(entry)
print('PASS: mechanism review and Jess response saved; manuscript unchanged.')
