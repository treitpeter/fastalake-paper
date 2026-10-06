"""Independent paired summaries and fixed-dictionary functional evidence; no annotation transfer."""
import csv,gzip,hashlib,json,os,re,sys
from pathlib import Path
from collections import defaultdict
import numpy as np
import pandas as pd
import ahocorasick
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
from search import peptides,sha,selection
B=Path(__file__).resolve().parents[1];m=json.loads((B/'manifest.json').read_text());pilot=len(sys.argv)>1 and sys.argv[1]=='pilot';O=B/('pilot_function_review' if pilot else 'figures');O.mkdir(exist_ok=False)
for rec in m['frozen_code']+m['annotation_inputs']:
 assert sha(rec['path'])==rec['sha256'],rec['path']
tasks=[t for t in m['tasks'] if not pilot or t['sample'] in m['pilot_samples']];sets={};tables=[];receipts=[]
for t in tasks:
 p=B/'results'/t['sample'];r=json.loads((p/'SEARCH_COMPLETE.json').read_text());assert r['status']=='PASS' and r['manifest_sha256']==sha(B/'manifest.json')
 for rec in r['outputs']:assert sha(rec['path'])==rec['sha256']
 receipts.append({'path':str(p/'SEARCH_COMPLETE.json'),'sha256':sha(p/'SEARCH_COMPLETE.json')});tab=pd.read_csv(p/'identification_summary.tsv',sep='\t');assert len(tab)==7;tables.append(tab)
 for arm in tab.arm:
  folder=Path(t['prior_search'])/arm if arm in ['unclustered','cluster_097','cluster_099'] else p/'searches'/arm
  peps=peptides(folder/'results.sage.tsv');assert len(peps)==int(tab.loc[tab.arm==arm,'peptides'].iloc[0]);sets[t['sample'],arm]=peps
 baseline=sets[t['sample'],'unclustered']
 for row in tab.itertuples():
  assert row.lost==len(baseline-sets[t['sample'],row.arm]) and row.gained==len(sets[t['sample'],row.arm]-baseline)
  assert abs(row.peptide_count_ratio-len(sets[t['sample'],row.arm])/len(baseline))<1e-10
all_peptides=set.union(*sets.values());ann=defaultdict(set);ann_seen=set();conflicts=defaultdict(set)
with gzip.open(m['annotations'],'rt') as f:
 for r in csv.DictReader(f,delimiter='\t'):
  h=r['sequence_sha256'];ks=set(re.findall(r'K\d{5}',r['KEGG_ko']));ann[h]|=ks;ann_seen.add(h);conflicts[h].add(tuple(sorted(ks)))
seqs={}
for rec in m['function_fastas']:
 assert sha(rec['path'])==rec['sha256']
 for _,seq in selection.fasta(rec['path']):seqs[hashlib.sha256(seq.encode()).hexdigest()]=seq
assert ann_seen<=set(seqs),'Annotation sequence hash absent from frozen reference pool'
auto=ahocorasick.Automaton()
for p in all_peptides:auto.add_word(p,p)
auto.make_automaton();pko=defaultdict(set);matched=set();annotated_match=set()
for h,seq in seqs.items():
 for _,pep in auto.iter(seq.replace('I','L')):
  matched.add(pep);pko[pep]|=ann[h]
  if h in ann_seen:annotated_match.add(pep)
with gzip.open(O/'fixed_peptide_KO_mapping.tsv.gz','wt') as f:
 w=csv.writer(f,delimiter='\t');w.writerow(['peptide','matches_fixed_sequence_pool','matches_annotated_sequence_record','compatible_KOs','KO_count'])
 for p in sorted(all_peptides):w.writerow([p,int(p in matched),int(p in annotated_match),';'.join(sorted(pko[p])),len(pko[p])])
ko_rows=[];ko_fates=[];member_rows=[]
def kos(peps,single=False):return set().union(*(pko[p] for p in peps if not single or len(pko[p])==1))
for t in tasks:
 sample=t['sample'];baseline=sets[sample,'unclustered'];bk=kos(baseline);bs=kos(baseline,True)
 with Path(t['members']).open() as f:members={r['peptide']:r['members'] for r in csv.DictReader(f,delimiter='\t') if r['n_members']=='1'}
 for arm in sorted(a for s,a in sets if s==sample):
  peps=sets[sample,arm];k=kos(peps);sk=kos(peps,True)
  ko_rows.append({'sample':sample,'arm':arm,'compatible_KOs':len(k),'baseline_compatible_KOs':len(bk),'baseline_KOs_retained':len(k&bk),'KO_recovery':len(k&bk)/len(bk) if bk else None,'KO_lost':len(bk-k),'KO_gained':len(k-bk),'single_KO_evidence_functions':len(sk),'baseline_single_KOs':len(bs),'single_KO_recovery':len(sk&bs)/len(bs) if bs else None,'peptides':len(peps),'unmapped_peptides':len(peps-matched),'mapped_without_KO':sum(p in matched and not pko[p] for p in peps),'single_KO_peptides':sum(len(pko[p])==1 for p in peps),'multi_KO_peptides':sum(len(pko[p])>1 for p in peps)})
  for ko in sorted(k|bk):ko_fates.append({'sample':sample,'arm':arm,'KO':ko,'baseline':int(ko in bk),'retained_arm':int(ko in k)})
  for member in sorted(set(members.values())):
   ref={p for p,v in members.items() if v==member};member_rows.append({'sample':sample,'arm':arm,'member':member,'member_specific_peptides':len(ref&peps),'baseline_member_specific_peptides':len(ref&baseline),'lost':len((ref&baseline)-peps),'gained':len((ref&peps)-baseline)})
df=pd.concat(tables,ignore_index=True);kf=pd.DataFrame(ko_rows)
df.to_csv(O/'per_acquisition_identifications.tsv',sep='\t',index=False);kf.to_csv(O/'per_acquisition_KO_coverage.tsv',sep='\t',index=False);pd.DataFrame(ko_fates).to_csv(O/'KO_fates.tsv.gz',sep='\t',index=False);pd.DataFrame(member_rows).to_csv(O/'member_specific_peptides.tsv',sep='\t',index=False)
summary=df.groupby('arm')[['protein_entry_ratio','peptide_count_ratio','peptide_reduction_pct','baseline_identity_recovery','REF_recovery']].median().join(kf.groupby('arm')[['KO_recovery','single_KO_recovery','KO_lost','KO_gained']].median());summary.to_csv(O/'paired_medians.tsv',sep='\t')
assert (kf[['unmapped_peptides','mapped_without_KO','single_KO_peptides','multi_KO_peptides']].sum(axis=1)==kf.peptides).all()
# Completion receipt records scientific scope even for pilot-only functional checks.
caption='Clustering acquisition-specific candidate proteins before parsimony: fixed MMseqs/Linclust settings, 80% bidirectional coverage, same Sage configuration and unclustered denominator for all thresholds. Points are acquisition-paired medians; full per-acquisition results and losses/gains accompany the figure. Functional coverage uses a fixed exact-sequence CAMPI FL+REF+MG annotation pool. Compatible KO evidence can support multiple functions; single-KO evidence is shown separately. Unmapped and unannotated peptides remain explicit. These are annotated-function evidence counts, not functional activity or proof of conserved function. The 50% threshold is a stress test; the experiment does not cluster the initial reservoir or calibrate whole-workflow FDR. Original S02 confidence drift remains recorded. Plot uses nominal min-seq-id settings; the pinned Linclust default identity estimator and other heuristics are unchanged.'
if not pilot:
 order=['unclustered','cluster_099','cluster_097','cluster_095','cluster_090','cluster_075','cluster_050'];x=np.arange(len(order));fig,ax=plt.subplots(1,2,figsize=(12,4.5));data=summary.loc[order]
 ax[0].plot(x,100*data.protein_entry_ratio,'o-',label='Protein entries');ax[0].plot(x,100*data.peptide_count_ratio,'o-',label='Peptide count');ax[0].plot(x,100*data.baseline_identity_recovery,'o-',label='Baseline peptide identities');ax[0].set_ylabel('Relative to unclustered baseline (%)')
 ax[1].plot(x,100*data.KO_recovery,'o-',label='Compatible KO coverage');ax[1].plot(x,100*data.single_KO_recovery,'o-',label='Single-KO peptide evidence');ax[1].set_ylabel('Baseline functional evidence retained (%)')
 for a in ax:a.set_xticks(x,['None','99','97','95','90','75','50']);a.set_xlabel('Clustering identity setting (%)');a.legend(frameon=False,fontsize=8);a.grid(alpha=.2)
 fig.tight_layout()
 for ext in ['png','pdf','svg']:fig.savefig(O/('CAMPI_clustering_function_benchmark.'+ext),dpi=220,bbox_inches='tight')
 doc=Document();doc.add_heading('CAMPI clustering and functional coverage — review proposal',0);doc.add_paragraph('New 95%,90%,75%,50% runs plus frozen97%/99% comparison. Awaiting visual and scientific review; not yet integrated into the manuscript.');doc.add_picture(str(O/'CAMPI_clustering_function_benchmark.png'),width=Inches(6.5));doc.add_paragraph(caption);doc.add_paragraph(summary.to_string());doc.save(O/'CAMPI_clustering_function_benchmark.docx');(O/'CAPTION.md').write_text(caption+'\n')
receipt={'status':'PASS','mode':'pilot' if pilot else 'full','acquisitions':len(tasks),'job_id':os.environ.get('SLURM_JOB_ID'),'manifest_sha256':sha(B/'manifest.json'),'source_receipts':receipts,'functional_dictionary_sequences':len(seqs),'sequence_records_with_annotations':len(ann_seen),'sequence_annotation_conflicts':sum(len(v)>1 for v in conflicts.values()),'scope':caption,'scientific_and_visual_review':'PENDING','outputs':[{'path':str(p),'sha256':sha(p)} for p in O.iterdir() if p.is_file()]};(O/'COMPLETE.json').write_text(json.dumps(receipt,indent=2)+'\n');print(receipt['mode'],'collection PASS',flush=True)
