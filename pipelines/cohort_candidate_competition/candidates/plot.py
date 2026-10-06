"""Plot per-acquisition Sage scored_candidates (FastaLake vs matched metagenome, paired rank-one spectra) against peptide gain from candidates/full/per_acquisition.tsv. Writes candidates/figures/."""
from pathlib import Path
import json,hashlib
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
H=Path(__file__).resolve().parent;D=H/'full';O=H/'figures';O.mkdir(exist_ok=True)
r=json.loads((D/'COMPLETE.json').read_text());assert r['status']=='PASS' and r['mode']=='full'
df=pd.read_csv(D/'per_acquisition.tsv',sep='\t');fig,axs=plt.subplots(2,2,figsize=(12,8));colors={'fl':'#D55E00','mg':'#0072B2'}
for col,c in enumerate(['PREDICT','CAPSCAN']):
 d=df[df.cohort==c].sort_values('sample');x=range(1,len(d)+1);ax=axs[0,col];right=ax.twinx()
 for a,label in [('fl','FastaLake'),('mg','Metagenome + human')]:
  ax.plot(x,d['median_candidates_'+a+'_paired'],color=colors[a],lw=1.1,label=label+' candidates');right.plot(x,d['peptide_ids_'+a],color=colors[a],ls=':',alpha=.5,lw=.7)
 ax.set_ylabel('Median scored candidates\nper paired reported spectrum');ax.set_yscale('log');right.set_ylabel('Peptide IDs in the acquisition');ax.set_xlabel('Acquisitions (fixed sample-ID order)');ax.set_title(c);ax.legend(frameon=False,fontsize=8)
 axs[1,col].scatter(d.median_paired_candidate_reduction_pct,d.peptide_gain_pct,s=12,alpha=.5,color='#555555');axs[1,col].axhline(0,color='grey',lw=.5);axs[1,col].axvline(0,color='grey',lw=.5);axs[1,col].set_xlabel('Median candidate reduction on paired spectra (%)');axs[1,col].set_ylabel('Peptide identification gain (%)')
fig.text(.5,.005,'Solid: scored candidates (left axis). Dotted: peptide IDs in that acquisition (right axis).',ha='center',fontsize=10)
fig.tight_layout(rect=(0,.03,1,1));files=[]
for ext in ['png','pdf','svg']:
 p=O/('candidate_competition_and_yield.'+ext);fig.savefig(p,dpi=220,bbox_inches='tight');files.append(p)
plt.close(fig)
caption='Current matched PREDICT and CapScan searches. Candidate counts are Sage 0.14.6 scored_candidates: candidates having at least one fragment match during preliminary scoring, summed across tested charge/isotope hypotheses. Counts are compared on identical reported rank-one spectrum IDs in both arms, including target and decoy winners before confidence filtering. Spectra unreported in either arm are excluded from paired candidate ratios and their counts are preserved in Source Data; they are not assigned zero candidates. Solid lines show median candidate counts on the paired spectra (left axis, log); dotted lines show actual canonical target peptide IDs per acquisition at peptide q<=0.01 (right axis). Bottom panels relate per-acquisition median paired candidate reductions to peptide yield differences; points are acquisitions with repeated specimens/participants, not independent biological replicates. Participant summaries are supplied separately. This quantifies associated search competition and identification yield, not a Bonferroni correction factor, isolated causal benefit, or whole-workflow FDR guarantee.'
(O/'CAPTION.md').write_text(caption+'\n');doc=Document();doc.add_heading('Candidate competition and identification yield',0);doc.add_paragraph('Figure proposal for Peter’s review.');doc.add_picture(str(files[0]),width=Inches(6.4));doc.add_paragraph(caption);p=O/'Candidate_competition_figure_proposal.docx';doc.save(p);files.append(p)
(O/'COMPLETE.json').write_text(json.dumps({'status':'FIGURES_GENERATED_PENDING_VISUAL_REVIEW','inputs':r,'files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in files}},indent=2)+'\n')
