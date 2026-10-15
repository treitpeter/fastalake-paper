"""Supplementary Table 5 in v42: all 726 shuffled-query acquisitions.

The directory name STab6 is retained for compatibility with historical versions.
Counts are observations per acquisition and seed, not globally unique peptides.
"""
import json
from pathlib import Path

import openpyxl
import pandas as pd

XLSX = "Supplementary_Table_5_v42.xlsx"
METHODS = ("bags", "combined", "exact", "mass")
READ_ME = [
    ["Scope", "All 726 acquisitions (CAMPI 15, healthy stool 58, CapScan 347, PREDICT 306), two seeds (1,452 acquisition-seed runs); q<=0.01; canonical peptide observations, not globally unique sequences"],
    ["Interpretation", "Diagnostic control acceptance; not calibrated production false-discovery proportion"],
]


def load(sd):
    sd = Path(sd)
    rows = pd.read_csv(sd / "null_receipt_rows_726.tsv", sep="\t")
    receipt = json.loads((sd / "collection_726.json").read_text())
    assert receipt["status"] == "DIAGNOSTICS_COMPLETE"
    assert receipt["completed"] == receipt["expected"] == 726 and not receipt["errors"]
    counts = rows.groupby(["cohort", "receipt"]).size()
    assert len(counts) == 726 and (counts == 80).all()
    mix = rows[["cohort", "receipt"]].drop_duplicates().groupby("cohort").size().to_dict()
    assert mix == receipt["completed_by_cohort"]
    assert set(rows.seed) == {1, 2} and set(rows.method) == set(METHODS)
    assert not rows.duplicated(["cohort", "receipt", "seed", "method", "q", "level"]).any()
    assert (rows.net == rows.accepted - rows.exact_accepted).all()
    assert (rows.net == rows.gained - rows.lost).all()
    assert (rows.groupby(["cohort", "receipt", "seed", "method"]).selected_proteins.nunique() == 1).all()
    acq = rows[rows.q.eq(.01) & rows.level.eq("canonical_peptide")].copy()
    assert len(acq) == 726 * 2 * 4
    summary = acq.groupby("method").agg(
        accepted=("accepted", "sum"),
        acquisition_seeds_with_acceptance=("accepted", lambda x: int((x > 0).sum())),
        selected_proteins=("selected_proteins", "sum"),
    ).loc[list(METHODS)]
    for method in METHODS:
        for col in ("accepted", "selected_proteins"):
            assert summary.loc[method, col] == receipt["roster_726"][col][method]
    return acq, summary


def numbers(sd):
    acq, summary = load(sd)
    values = {"acquisitions": acq.receipt.nunique(), "seeds": acq.seed.nunique(), "q": .01}
    for method, row in summary.iterrows():
        values.update({method + "_" + col: int(value) for col, value in row.items()})
    return values


def write(sd, out):
    acq, summary = load(sd)
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    acq.to_csv(out / "null_control_by_acquisition_seed_v42.tsv", sep="\t", index=False)
    summary.to_csv(out / "null_control_summary_v42.tsv", sep="\t")
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Read me"
    for row in READ_ME:
        ws.append(row)
    ws = book.create_sheet("Summary")
    ws.append(["method", *summary.columns])
    for row in summary.reset_index().itertuples(index=False, name=None):
        ws.append(row)
    ws = book.create_sheet("Acquisition_seed_rows")
    ws.append(list(acq.columns))
    for row in acq.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v for v in row])
    book.save(out / XLSX)
    return out / XLSX
