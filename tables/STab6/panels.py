"""Supplementary Table 6: shuffled-query controls, 60 acquisitions x 2 seeds.

Rebuilds the three sheets of ``Supplementary_Table_6.xlsx`` (``Read me``, ``Summary``,
``Acquisition_seed_rows``) and the two support tables
(``null_control_summary.tsv``, ``null_control_by_acquisition_seed.tsv``) from
``source_data/STab6/``:

* ``null_receipt_rows.tsv``     all 80 rows of each of the 60 frozen receipts (input)
* ``null_receipts.tsv``         one row per receipt: status flags (checks only)
* ``null_receipt_results.tsv``  sha256 of every Sage result each receipt read (checks only)

The table logic is moved unchanged from the two original producer steps:

* ``package_release.py`` (lines 20-26): keep rows with ``q == 0.01`` and
  ``level == 'canonical_peptide'``, write them as the acquisition-seed table, sum
  ``accepted`` and ``selected_proteins`` by ``method`` for the summary, and assert
  ``mass == 13512`` and ``exact == 0``.
* ``package_v15.py`` (lines 54-58): write the workbook with a two-row ``Read me`` sheet
  and the two TSVs read back with ``csv.reader`` (so every cell is a string, as published).

There is no statistical test in this table. Counts are sums over acquisition-seed runs.
"""

import csv
from pathlib import Path

import openpyxl
import pandas as pd

Q_CUTOFF = 0.01                     # peptide q threshold; v28d line 485 ("At peptide q≤0.01")
LEVEL = "canonical_peptide"         # count level; v28d line 485 ("canonical peptide–acquisition–seed observations")
METHODS = ("exact", "bags", "mass", "combined")   # lookup routes, in the order the legend prints them (line 485)
N_RECEIPTS_FROZEN = 60              # receipts in the frozen snapshot; v28d line 485 ("60 acquisitions")
SEEDS = (1, 2)                      # shuffle seeds; v28d line 485 ("two seeds")
ROWS_PER_RECEIPT = 80               # 2 seeds x 5 q thresholds x 2 levels x 4 methods (producer finish.py assertion)

# Text of the "Read me" sheet, exactly as written by package_v15.py line 54.
README_ROWS = [
    ["Scope", "Frozen 60 acquisition receipts, two seeds; q<=0.01; canonical peptide observations, "
              "not globally unique sequences"],
    ["Interpretation", "Diagnostic control acceptance; not calibrated production false-discovery proportion"],
]

TSV_BY_ACQ = "null_control_by_acquisition_seed.tsv"
TSV_SUMMARY = "null_control_summary.tsv"
XLSX = "Supplementary_Table_6.xlsx"

_WORDS = {0: "zero", 1: "one", 2: "two", 3: "three", 4: "four", 5: "five"}


def load(sd):
    """Read the source tables.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab6``.

    Returns
    -------
    dict of pandas.DataFrame
        Keys ``rows``, ``receipts``, ``results``.
    """
    sd = Path(sd)
    return {
        "rows": pd.read_csv(sd / "null_receipt_rows.tsv", sep="\t"),
        "receipts": pd.read_csv(sd / "null_receipts.tsv", sep="\t"),
        "results": pd.read_csv(sd / "null_receipt_results.tsv", sep="\t"),
    }


def build_acquisition_seed_rows(d):
    """Acquisition-seed rows at the table's filter (producer ``package_release.py`` line 23).

    Parameters
    ----------
    d : dict
        Output of :func:`load`.

    Returns
    -------
    pandas.DataFrame
        One row per receipt (acquisition), seed and method; 60 x 2 x 4 = 480 rows.
    """
    r = d["rows"]
    return r[(r["q"] == Q_CUTOFF) & (r["level"] == LEVEL)].reset_index(drop=True)


def build_summary(acq):
    """Sum accepted observations and selected protein entries by route (``package_release.py`` line 25).

    Parameters
    ----------
    acq : pandas.DataFrame
        Output of :func:`build_acquisition_seed_rows`.

    Returns
    -------
    pandas.DataFrame
        Index ``method``; columns ``accepted``, ``selected_proteins``.
    """
    return acq.groupby("method")[["accepted", "selected_proteins"]].sum()


def write_table(acq, summary, out):
    """Write the two TSVs and the workbook exactly as the producers did.

    Parameters
    ----------
    acq, summary : pandas.DataFrame
        Outputs of :func:`build_acquisition_seed_rows` and :func:`build_summary`.
    out : pathlib.Path
        Output directory.

    Returns
    -------
    dict of pathlib.Path
        Keys ``by_acq``, ``summary``, ``xlsx``.
    """
    out = Path(out)
    p_acq, p_sum, p_x = out / TSV_BY_ACQ, out / TSV_SUMMARY, out / XLSX
    acq.to_csv(p_acq, sep="\t", index=False)
    summary.to_csv(p_sum, sep="\t")
    book = openpyxl.Workbook()
    read = book.active
    read.title = "Read me"
    for row in README_ROWS:
        read.append(row)
    for sheet, file in [("Summary", p_sum), ("Acquisition_seed_rows", p_acq)]:
        ws = book.create_sheet(sheet)
        with open(file, newline="") as fh:
            for row in csv.reader(fh, delimiter="\t"):
                ws.append(row)
    book.save(p_x)
    return {"by_acq": p_acq, "summary": p_sum, "xlsx": p_x}


def compare_xlsx_values(new, expected):
    """Compare two workbooks sheet by sheet and cell by cell (values only).

    Workbook bytes carry a save timestamp, so byte identity is not expected.

    Returns
    -------
    dict
        ``identical`` (bool), ``sheets``, and up to 20 ``differences``.
    """
    a, b = openpyxl.load_workbook(new), openpyxl.load_workbook(expected)
    diffs = []
    if a.sheetnames != b.sheetnames:
        diffs.append(f"sheet names {a.sheetnames} vs {b.sheetnames}")
    for name in b.sheetnames:
        if name not in a.sheetnames:
            continue
        va = list(a[name].iter_rows(values_only=True))
        vb = list(b[name].iter_rows(values_only=True))
        if len(va) != len(vb):
            diffs.append(f"{name}: {len(va)} vs {len(vb)} rows")
        for i, (x, y) in enumerate(zip(va, vb), 1):
            if x != y:
                diffs.append(f"{name} row {i}: {x} vs {y}")
    return {"identical": not diffs, "sheets": b.sheetnames, "differences": diffs[:20]}


def checks(d):
    """Repeat the producers' checks on the source tables.

    Checks: 60 receipts, 80 rows each, both seeds and all four methods present; every
    receipt has status NULL_DIAGNOSTIC_COMPLETE_REVIEW_REQUIRED and workflow_fdr
    UNVERIFIED; ``net == accepted - exact_accepted == gained - lost`` on every row
    (``finish.py``); ``selected_proteins`` is constant across q and level within a
    receipt, seed and method; every receipt lists 12 Sage result files with sha256; the
    ``package_release.py`` assertions ``mass == 13512`` and ``exact == 0``.

    Returns
    -------
    tuple
        ``(pandas.DataFrame of check, value, pass)``, ``all_pass``.
    """
    r, rec, res = d["rows"], d["receipts"], d["results"]
    acq = build_acquisition_seed_rows(d)
    s = build_summary(acq)
    per = r.groupby(["receipt", "cohort"]).size()
    sp = r.groupby(["receipt", "seed", "method"])["selected_proteins"].nunique()
    nres = res.groupby("receipt").size()
    out = [
        ("receipts", len(per), len(per) == N_RECEIPTS_FROZEN),
        ("rows per receipt (min,max)", (int(per.min()), int(per.max())),
         per.min() == per.max() == ROWS_PER_RECEIPT),
        ("seeds", sorted(r["seed"].unique().tolist()), tuple(sorted(r["seed"].unique())) == SEEDS),
        ("methods", sorted(r["method"].unique().tolist()), set(r["method"]) == set(METHODS)),
        ("status all NULL_DIAGNOSTIC_COMPLETE_REVIEW_REQUIRED",
         int((rec["status"] == "NULL_DIAGNOSTIC_COMPLETE_REVIEW_REQUIRED").sum()),
         bool((rec["status"] == "NULL_DIAGNOSTIC_COMPLETE_REVIEW_REQUIRED").all())),
        ("workflow_fdr all UNVERIFIED", int((rec["workflow_fdr"] == "UNVERIFIED").sum()),
         bool((rec["workflow_fdr"] == "UNVERIFIED").all())),
        ("net == accepted - exact_accepted", int((r["net"] == r["accepted"] - r["exact_accepted"]).sum()),
         bool((r["net"] == r["accepted"] - r["exact_accepted"]).all())),
        ("net == gained - lost", int((r["net"] == r["gained"] - r["lost"]).sum()),
         bool((r["net"] == r["gained"] - r["lost"]).all())),
        ("selected_proteins constant over q and level", int(sp.max()), int(sp.max()) == 1),
        ("Sage result files per receipt (min,max)", (int(nres.min()), int(nres.max())),
         nres.min() == nres.max() == 12 and len(nres) == N_RECEIPTS_FROZEN),
        ("acquisition-seed rows at filter", len(acq), len(acq) == N_RECEIPTS_FROZEN * len(SEEDS) * len(METHODS)),
        ("producer assert mass == 13512", int(s.loc["mass", "accepted"]), int(s.loc["mass", "accepted"]) == 13512),
        ("producer assert exact == 0", int(s.loc["exact", "accepted"]), int(s.loc["exact", "accepted"]) == 0),
    ]
    df = pd.DataFrame(out, columns=["check", "value", "pass"])
    return df, bool(df["pass"].all())


def word(n):
    """Spell a small count as the manuscript does ("zero", "two")."""
    return _WORDS[int(n)]


def stat_n_acquisitions(d):
    """Number of distinct acquisitions (receipts) in the table.

    Unit: acquisition; n = 60 (15 CAMPI, 43 HSTOOL, 1 CAPSCAN, 1 PREDICT).

    Returns
    -------
    int
    """
    return int(build_acquisition_seed_rows(d)["receipt"].nunique())


def stat_n_seeds(d):
    """Number of shuffle seeds in the table (unit: seed per acquisition; n = 60 acquisitions).

    Returns
    -------
    int
    """
    return int(build_acquisition_seed_rows(d)["seed"].nunique())


def stat_q_cutoff(d):
    """The single peptide q threshold of every row in the table (unit: acquisition-seed run; n = 480 rows).

    Returns
    -------
    float
    """
    q = build_acquisition_seed_rows(d)["q"].unique()
    assert len(q) == 1
    return float(q[0])


def stat_accepted(d, method):
    """Accepted canonical peptide-acquisition-seed observations for one route at q <= 0.01.

    Unit: peptide-acquisition-seed observation, summed over 60 acquisitions x 2 seeds
    (120 runs per route); not globally distinct sequences.

    Returns
    -------
    int
    """
    return int(build_summary(build_acquisition_seed_rows(d)).loc[method, "accepted"])


def stat_selected_proteins(d, method):
    """Selected protein entries for one route, summed over 120 acquisition-seed runs.

    Unit: protein entry per acquisition-seed run; not deduplicated across runs.

    Returns
    -------
    int
    """
    return int(build_summary(build_acquisition_seed_rows(d)).loc[method, "selected_proteins"])


def cohort_mix(d):
    """Acquisitions per cohort in the frozen snapshot (unit: acquisition; n = 60).

    Returns
    -------
    pandas.Series
    """
    return build_acquisition_seed_rows(d).drop_duplicates("receipt")["cohort"].value_counts()
