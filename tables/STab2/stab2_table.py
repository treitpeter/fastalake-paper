"""Supplementary Table 2: full-route shuffled-control acceptance and control exposure.

Rebuilds the published workbook ``Supplementary_Table_2.xlsx`` (18 sheets) from
``source_data/STab2/`` and recomputes the two counts that the Methods cite to this table
(v28d line 743).

Sheets and the source_data files they come from:

=====================================  ==============================================
sheet                                  source_data/STab2/
=====================================  ==============================================
Read_me                                constant text (``README_ROWS``)
Cohort_overview                        computed by ``cohort_overview`` from
                                       per_acquisition.tsv, stage_control_counts.tsv,
                                       baseline_comparison.tsv
Per_acquisition                        per_acquisition.tsv
Stage_controls                         stage_control_counts.tsv
Baseline_comparison                    baseline_comparison.tsv
Quantile_summary                       quantile_summary.tsv (recomputed and checked by
                                       ``quantile_summary``)
Control_discoveries                    control_discoveries_<cohort>.tsv
fusion_Read_me .. fusion_All_acquisitions  fusion_*.tsv, fusion_validation.json
indep_Read_me .. indep_Fusion_all_thresholds  indep_*.tsv, indep_validation.json
=====================================  ==============================================

The code is moved from the original producers with paths made repository-relative:
``build_package.py`` (2026-09-06; Cohort_overview logic and the first seven sheets),
``collect_cohort.py`` (quartile summary), ``package_streamlining_sources.py`` and
``package_v10_sources.py`` (fusion and independent-selection notes), and
``package_edit3.py`` (merging those two workbooks as prefixed sheets). The later
path-scrubbing step (``package_v15.py``) changed no cell value in this workbook.
"""
import json
import statistics
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

COHORTS = ["CAMPI_MOCK", "HSTOOL", "PREDICT", "CAPSCAN"]  # build_package.py COHORTS
LABELS = ["CAMPI", "Healthy stool", "PREDICT", "CapScan"]  # build_package.py LABELS
PRIMARY_Q = 0.01            # peptide q threshold of the primary view; Read_me row "Primary"
PRIMARY_MIN_RESIDUES = 5    # actual peptide length floor of the primary view (5-50)
N_ACQUISITIONS = 710        # 15 CAMPI + 58 healthy stool + 306 PREDICT + 331 CapScan
CONTROL_SEED = 20260906     # shuffled-control construction seed (experiment.json control_source_stats.seed)
QUANTILE_METHOD = "exclusive"  # statistics.quantiles method in collect_cohort.py
FUSION_SEEDS = (20260907, 20260908)  # fusion control seeds (fusion sheets, column seed)
POOL_ARMS = {"own": "acquisition-specific databases", "pool_other11": "other-eleven unions"}  # v28d line 743

#: Read_me sheet, verbatim from build_package.py.
README_ROWS = [
    ["Population", "710 acquisitions; one joint control-bank augmentation per cohort"],
    ["Discovery unit", "Modification-stripped, I/L-collapsed peptide sequence; label 1"],
    ["Primary", "peptide_q <= 0.01, actual lengths 5–50"],
    ["Sensitivity", "Actual lengths 9–50; q thresholds 0.001, 0.005, 0.02, 0.05"],
    ["Undefined", "An empty discovery set has an undefined fraction, not zero"],
    ["Original overlap", "Exact I/L-normalised substring anywhere in the original stage-0 reservoir"],
    ["Interpretation", "Observed control-labelled fractions, not a paired or scaled total FDR estimate"],
    ["Dispersion", "Acquisition quartiles; not independent-participant confidence intervals"],
    ["Pooling", "15/58/306/331; all 710 ordinary-target candidate FASTAs are byte-identical to archived extraction despite changed shared pool sizes"],
    ["Synthetic bank", "2,333,639 controls; approximately 1% hash sample of the original CAMPI reservoir"],
    ["Foreign bank", "512 P. furiosus plus 16,418 A. thaliana; not assumed absent"],
]

#: fusion_Read_me rows, verbatim from package_streamlining_sources.py; "Validation" is the JSON text.
FUSION_NOTES = [
    ("Figure", "SupplementaryFigure6c–f; 1,420 completed searches, two seeds"),
    ("Inference", "Assumption-dependent diagnostics; whole-procedure FDR control is not established"),
    ("Aggregation", "Ratios of summed acquisition counts; repeated sequences across acquisitions remain repeated"),
    ("Validation", "fusion_validation.json"),
    ("Archive", "Frozen completed per-acquisition table and audit records accompany workbook"),
]
FUSION_FRAMES = [  # (sheet name before the 31-character cut, source_data file)
    ("cohort_q001", "fusion_cohort_q001.tsv"),
    ("cohort_threshold_summary", "fusion_cohort_threshold_summary.tsv"),
    ("matched_seed_summary", "fusion_matched_seed_summary.tsv"),
    ("short_5_to_8_summary", "fusion_short_5_to_8_summary.tsv"),
    ("acquisition_outliers_q001", "fusion_acquisition_outliers_q001.tsv"),
    ("completeness", "fusion_completeness.tsv"),
    ("All_acquisitions", "fusion_all_acquisitions.tsv"),
]

#: indep_Read_me rows, verbatim from package_v10_sources.py.
INDEP_NOTES = [
    ("Design", "S01/S02 fixed/self/independent; independent selection excludes evaluated spectra at cohort, acquisition and razor stages."),
    ("Replication", "Two technical acquisitions of one community. Seeds are control-realisation sensitivities, not biological replicates."),
    ("Reference", "Same 29,673-sequence deposited reference; substantially smaller than production reservoir."),
    ("Interpretation", "Identification yields, identity intersections and assumption-dependent diagnostics. Not generating-spectrum truth or whole-procedure calibration."),
    ("Validation", "indep_validation.json"),
]
INDEP_FRAMES = [
    ("Ordinary_counts", "indep_ordinary_counts.tsv"),
    ("Fusion_all_thresholds", "indep_fusion_all_thresholds.tsv"),
]


def _tsv(sd, name):
    return pd.read_csv(Path(sd) / name, sep="\t")


def load(sd):
    """Read the full-route tables.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab2``.

    Returns
    -------
    dict of pandas.DataFrame
        ``per_acquisition`` (7,100 rows: 710 acquisitions x 5 q thresholds x 2 length
        floors), ``stage_control_counts``, ``baseline_comparison`` (710 rows),
        ``summary`` and ``control_discoveries`` (all four cohorts, cohort order).
    """
    t = {n: _tsv(sd, n + ".tsv") for n in ["per_acquisition", "stage_control_counts", "baseline_comparison"]}
    t["summary"] = _tsv(sd, "quantile_summary.tsv")
    t["control_discoveries"] = pd.concat([_tsv(sd, f"control_discoveries_{c}.tsv") for c in COHORTS],
                                         ignore_index=True)
    assert len(t["baseline_comparison"]) == N_ACQUISITIONS
    assert not t["baseline_comparison"][["cohort", "sample"]].duplicated().any()
    assert len(t["per_acquisition"]) == N_ACQUISITIONS * 5 * 2
    return t


def cohort_overview(data, stages, comparison):
    """Cohort_overview sheet; logic of build_package.py unchanged.

    Parameters
    ----------
    data, stages, comparison : pandas.DataFrame
        Per_acquisition, Stage_controls and Baseline_comparison tables.

    Returns
    -------
    pandas.DataFrame
        One row per cohort. Medians and counts are across acquisitions (statistical unit:
        acquisition; n = 15, 58, 306, 331) at peptide q <= 0.01 and lengths 5-50.
    """
    primary = data[(data.q_threshold == PRIMARY_Q) & (data.min_residues == PRIMARY_MIN_RESIDUES)]
    compact = []
    for c, label in zip(COHORTS, LABELS):
        x = primary[primary.cohort == c]
        trace = stages[(stages.cohort == c) & (stages.control_class == "shuffle")]
        comp = comparison[comparison.cohort == c]
        compact.append(dict(
            cohort=c, label=label, n=len(x), n_defined=int(x.shuffle_only_fraction.notna().sum()),
            evidence_shuffle=int(trace[trace.stage == "1_cohort"].n_proteins.iloc[0]),
            acquisition_shuffle_median=float(trace[trace.stage == "3_acquisition"].n_proteins.median()),
            razor_shuffle_median=float(trace[trace.stage == "4_razor"].n_proteins.median()),
            accepted_shuffle_median=float(x.n_shuffle_only.median()),
            accepted_shuffle_not_original_median=float(x.shuffle_only_not_in_original_reservoir.median()),
            shuffle_fraction_median_pct=float(100 * x.shuffle_only_fraction.median()),
            shuffle_not_original_median_pct=float(100 * x.shuffle_only_not_in_original_reservoir_fraction.median()),
            acquisitions_with_shuffle=int((x.n_shuffle_only > 0).sum()),
            acquisitions_with_shuffle_not_original=int((x.shuffle_only_not_in_original_reservoir > 0).sum()),
            plant_not_original_median_pct=float(100 * x.plant_only_not_in_original_reservoir_fraction.median()),
            thermophilic_not_original_median_pct=float(100 * x.thermophilic_only_not_in_original_reservoir_fraction.median()),
            target_count_retention_median_pct=float(100 * comp.count_ratio.median()),
            target_set_recovery_median_pct=float(100 * comp.own_set_recovery.median())))
    return pd.DataFrame(compact)


def quantile_summary(sd):
    """Recompute the Quantile_summary sheet; logic of collect_cohort.py unchanged.

    Reads per_acquisition.tsv as text, as collect_cohort.py read its own rows, so values
    are parsed exactly as the original step parsed them.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab2``.

    Returns
    -------
    pandas.DataFrame
        Per cohort, q threshold, length floor, control class and metric: median, first and
        third quartile (``statistics.quantiles``, method ``exclusive``) and maximum of the
        per-acquisition fraction (statistical unit: acquisition; undefined fractions of
        empty discovery sets are excluded and counted in ``n_defined``).
    """
    import csv
    import collections
    rows = list(csv.DictReader(open(Path(sd) / "per_acquisition.tsv"), delimiter="\t"))
    medians = []
    for c in COHORTS:
        groups = collections.defaultdict(list)
        for r in rows:
            if r["cohort"] == c:
                groups[float(r["q_threshold"]), int(r["min_residues"])].append(r)
        for (q, floor), rr in sorted(groups.items()):
            for cls in ["shuffle", "thermophilic", "plant"]:
                for metric in [cls + "_only_fraction", cls + "_only_not_in_original_reservoir_fraction"]:
                    v = [float(r[metric]) for r in rr if r[metric] not in ["", None]]
                    qu = statistics.quantiles(v, method=QUANTILE_METHOD) if len(v) > 1 else [None] * 3
                    medians.append(dict(cohort=c, q_threshold=q, min_residues=floor, control_class=cls,
                                        metric=metric, n_attempted=len(rr), n_defined=len(v),
                                        median=statistics.median(v) if v else None, q1=qu[0], q3=qu[2],
                                        max=max(v) if v else None))
    return pd.DataFrame(medians)


def _notes_rows(notes, sd):
    out = [["topic", "definition"]]
    for k, v in notes:
        out.append([k, (Path(sd) / v).read_text() if v.endswith(".json") else v])
    return out


def _frame_rows(df):
    rows = [list(df.columns)]
    for row in df.itertuples(index=False, name=None):
        rows.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in row])
    return rows


def sheets(sd):
    """All 18 sheets in published order.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab2``.

    Returns
    -------
    list of (str, list of list)
        Sheet name and its rows (header first).
    """
    t = load(sd)
    out = [("Read_me", [list(r) for r in README_ROWS])]
    for name, d in [("Cohort_overview", cohort_overview(t["per_acquisition"], t["stage_control_counts"], t["baseline_comparison"])),
                    ("Per_acquisition", t["per_acquisition"]), ("Stage_controls", t["stage_control_counts"]),
                    ("Baseline_comparison", t["baseline_comparison"]), ("Quantile_summary", t["summary"]),
                    ("Control_discoveries", t["control_discoveries"])]:
        out.append((name, _frame_rows(d)))
    for prefix, notes, frames in [("fusion", FUSION_NOTES, FUSION_FRAMES), ("indep", INDEP_NOTES, INDEP_FRAMES)]:
        out.append(((prefix + "_Read_me")[:31], _notes_rows(notes, sd)))
        for name, f in frames:
            out.append(((prefix + "_" + name)[:31], _frame_rows(_tsv(sd, f))))
    return out


def write_workbook(sd, path):
    """Write the rebuilt workbook.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab2``.
    path : pathlib.Path
        Output ``.xlsx``.

    Returns
    -------
    list of (str, int)
        Sheet names and row counts written.
    """
    wb = Workbook(write_only=True)
    written = []
    for name, rows in sheets(sd):
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        written.append((name, len(rows)))
    wb.save(path)
    return written


def _same(a, b):
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b and type(a) is type(b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return float(a) == float(b)
    return a == b


def compare_workbooks(built, expected):
    """Compare every cell of two workbooks.

    Parameters
    ----------
    built, expected : pathlib.Path
        Rebuilt and published ``.xlsx``.

    Returns
    -------
    dict
        Per sheet: rows in each, number of differing cells and up to five examples;
        ``sheet_order_equal`` and ``identical`` (all values equal, numbers compared
        exactly as floats).
    """
    a = load_workbook(built, read_only=True)
    b = load_workbook(expected, read_only=True)
    res = {"sheet_order_equal": a.sheetnames == b.sheetnames, "sheets": {}}
    for name in b.sheetnames:
        if name not in a.sheetnames:
            res["sheets"][name] = {"missing_in_rebuild": True}
            continue
        ra = [tuple(r) for r in a[name].iter_rows(values_only=True)]
        rb = [tuple(r) for r in b[name].iter_rows(values_only=True)]
        diffs, ex = 0, []
        for i in range(max(len(ra), len(rb))):
            x = ra[i] if i < len(ra) else ()
            y = rb[i] if i < len(rb) else ()
            for j in range(max(len(x), len(y))):
                u = x[j] if j < len(x) else None
                v = y[j] if j < len(y) else None
                if not _same(u, v):
                    diffs += 1
                    if len(ex) < 5:
                        ex.append({"row": i + 1, "col": j + 1, "rebuilt": repr(u), "published": repr(v)})
        res["sheets"][name] = {"rows_rebuilt": len(ra), "rows_published": len(rb), "differing_cells": diffs, "examples": ex}
    res["identical"] = res["sheet_order_equal"] and all(s.get("differing_cells", 1) == 0 for s in res["sheets"].values())
    return res


def stat_line743(sd):
    """Shuffle-only accepted peptide-acquisition observations, healthy-stool pooling experiment.

    Twelve healthy-stool technical-replicate acquisitions (statistical unit:
    peptide-acquisition observation; n = 12 acquisitions per arm), peptide q <= 0.01,
    lengths 5-50. Numerator: accepted peptides assigned only to shuffled controls and
    absent from the original reservoir; denominator: all accepted canonical peptides.
    Counts are summed over acquisitions, so a sequence found in two acquisitions counts
    twice.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab2``.

    Returns
    -------
    dict
        For arms ``own`` and ``pool_other11``: ``n_acquisitions``, ``shuffle_only`` and
        ``accepted``.
    """
    d = _tsv(sd, "hstool_pooling_per_acquisition.tsv")
    p = d[(d.q_threshold == PRIMARY_Q) & (d.min_residues == PRIMARY_MIN_RESIDUES)]
    out = {}
    for arm in POOL_ARMS:
        x = p[p.arm == arm]
        out[arm] = {"n_acquisitions": int(len(x)),
                    "shuffle_only": int(x.shuffle_only_not_in_original_reservoir.sum()),
                    "accepted": int(x.n_canonical.sum())}
    return out


def find_in_workbook(path, values):
    """Report whether any cell of a workbook equals one of ``values``.

    Used to test whether the counts cited to Supplementary Table 2 are in it.

    Parameters
    ----------
    path : pathlib.Path
        ``.xlsx`` to search.
    values : iterable of int
        Values to look for (numbers, or strings with thousands separators).

    Returns
    -------
    dict
        value -> list of ``sheet!row,col`` hits (at most five each).
    """
    targets = {int(v) for v in values}
    strs = {f"{v:,}" for v in targets} | {str(v) for v in targets}
    hits = {v: [] for v in targets}
    wb = load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            for j, c in enumerate(row, 1):
                if isinstance(c, (int, float)) and not isinstance(c, bool) and c in targets and len(hits[int(c)]) < 5:
                    hits[int(c)].append(f"{ws.title}!{i},{j}")
                elif isinstance(c, str) and any(s in c for s in strs):
                    for v in targets:
                        if (f"{v:,}" in c) and len(hits[v]) < 5:
                            hits[v].append(f"{ws.title}!{i},{j} (text)")
    return hits
