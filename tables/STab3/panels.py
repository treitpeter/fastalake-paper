"""Supplementary Table 3: PREDICT antibiotic endpoints under alternative handling of
rifaximin for hepatic encephalopathy (HE).

The workbook has three sheets: ``Read_me``, ``Endpoints`` (two endpoints for each of four
exposure definitions) and ``Composition`` (mean origin composition by state and definition).

Two routes:

* Public: :func:`build_workbook` rebuilds the published workbook from
  ``source_data/STab3/endpoints_by_definition.tsv`` and
  ``source_data/STab3/composition_by_definition.tsv`` (port of the Source Data builder,
  ``specials.py::st6``), and :func:`compare_workbooks` checks it cell by cell against
  ``expected/Supplementary_Table_3.xlsx``.
* Controlled: :func:`stat_definitions` recomputes every statistic and count from the two
  controlled inputs (port of ``iteration1/fix_rifaximin/rifaximin_sensitivity.py``; the
  state rule, aggregation, bootstrap, sign-flip test and Benjamini-Hochberg step are moved
  unchanged). It needs ``FASTALAKE_CONTROLLED_DIR`` with
  ``PREDICT_metadata_2026-08-03.tsv`` and ``acquisition_accounting.tsv``.

Definitions (``scenario`` column)

A_current
    Primary rule; the field ``visit_rifaximin_for_he`` is not used.
B_HE_as_exposure
    ``visit_rifaximin_for_he`` added to the core flags (1 makes an acquisition On; Off needs it 0).
C_HE_patients_excluded
    Patients with any acquisition where ``visit_rifaximin_for_he`` is 1 are removed.
D_HE_off_rows_excluded
    Only acquisitions that are Off under A and have ``visit_rifaximin_for_he`` 1 set to Unknown.

The upstream excerpt of the original script that counted the EDS 85-patient display under
A and B is not ported: it does not enter this table (see PROVENANCE.md).
"""
import hashlib
import io
import os
import re
from pathlib import Path

import numpy as np
import pandas as pd

# ---------------------------------------------------------------------------------
# Choices (named constants; see README.md, Choices)
# ---------------------------------------------------------------------------------
SEED = 20260907                 # Methods, v28d line 795 (bootstrap); sign flips use SEED + endpoint index
N_BOOTSTRAP = 50_000            # Methods, v28d line 795 (patient resamples)
N_SIGNFLIP_BATCHES = 100        # Methods, v28d line 795: 100 x 2,000 = 200,000 sign assignments
SIGNFLIP_BATCH = 2_000
CI_QUANTILES = (0.025, 0.975)   # 95% percentile interval, numpy.quantile default (linear)
TIE_TOLERANCE = 1e-12           # |null mean| >= |observed mean| - 1e-12 counts as extreme
PSEUDOCOUNT = 1                 # plus-one correction, Methods v28d line 795
ENDPOINTS = [("Human", 100), ("pg_parsimony", 1)]   # (column, scale); Human -> percentage points
CLASSES = ['Human', 'Microbial_species', 'Microbial_above_species', 'Other_annotated_origins', 'Unresolved_origin']
CORE_FLAGS = ['visit_antibiotic_prophylaxis_for_sbp', 'visit_antibiotic_prophylaxis_for_variceal_bleeding',
              'visit_antibiotic_prophylaxis_for_intervention', 'visit_other_type_of_antibiotic',
              'visit_antibiotic_for_infection_nonprophylactic', 'visit_rifaximin', 'visit_norfloxacin',
              'visit_cephalosporines', 'visit_penicillines']   # Methods, v28d line 793
HE_FIELD = 'visit_rifaximin_for_he'
N_ACQUISITIONS = 306
SCENARIOS = ['A_current', 'B_HE_as_exposure', 'C_HE_patients_excluded', 'D_HE_off_rows_excluded']

#: Controlled inputs and the sha256 of the exact files the published values came from.
CONTROLLED = {
    'PREDICT_metadata_2026-08-03.tsv': 'fc5dbf07df2fd25c7def3e9d583b2745c9f237b12a6a9c6e78ecefb249e90238',
    'acquisition_accounting.tsv': '3880642bef93ccb390ff7be86b30a4aec715134b4f96c63ce0d5c6a4f5deaa23',
}

#: Read_me sheet text, verbatim from the Source Data builder (specials.py ST6_README).
README_LINES = [
    'Supplementary Table 3. PREDICT antibiotic endpoints under alternative handling of rifaximin for hepatic encephalopathy.',
    '',
    'Rows give the primary definition (42 patients), rifaximin for hepatic encephalopathy counted as exposure (40 patients), patients with any such record excluded (35 patients) and only off-state acquisitions with such a record excluded (38 patients).',
    'Endpoints: for human-assigned share of total LFQ intensity (percentage points) and quantified protein groups, the number of paired patients and acquisitions, state means, mean and median paired differences (on minus off), number of patients with a higher value on antibiotics, 95% bootstrap interval of the mean difference, sign-flip P value and Benjamini-Hochberg q value across the two endpoints within each definition.',
    'Composition: mean origin composition of total LFQ intensity (fractions) by antibiotic state for each definition.',
    'All definitions used the same resampling settings and seeds as the primary analysis.',
]

NUMBER_WORDS = {2: 'two', 7: 'seven', 12: 'twelve'}


def sha(p):
    """Return the sha256 hex digest of file ``p``."""
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ---------------------------------------------------------------------------------
# Public route: the workbook
# ---------------------------------------------------------------------------------
def _tsv_rows(path):
    """Read a TSV as a list of string rows (as the Source Data builder did)."""
    import csv
    with open(path, newline='', encoding='utf-8') as fh:
        return list(csv.reader(fh, delimiter='\t'))


def build_workbook(sd, out_path):
    """Rebuild Supplementary Table 3 from source_data (port of specials.py::st6).

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab3``.
    out_path : pathlib.Path
        Where to write the ``.xlsx``.

    Returns
    -------
    pathlib.Path
        ``out_path``.
    """
    import openpyxl
    rows = _tsv_rows(sd / 'endpoints_by_definition.tsv')
    comp = _tsv_rows(sd / 'composition_by_definition.tsv')
    # source_data writes the state column as 'state' (release scan); the sheet header is restored
    assert comp[0][1] == 'state'
    comp[0][1] = 'antibiotic_state'
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'Read_me'
    for line in README_LINES: ws.append([line])
    for title, tab in (('Endpoints', rows), ('Composition', comp)):
        w = wb.create_sheet(title)
        for i, r in enumerate(tab):
            if i == 0:
                w.append(r); continue
            conv = []
            for x in r:
                try:
                    conv.append(int(x) if re.fullmatch(r'-?\d+', x) else float(x))
                except ValueError:
                    conv.append(x)
            w.append(conv)
    wb.properties.creator = 'FastaLake'
    out_path = Path(out_path)
    wb.save(out_path)
    return out_path


def compare_workbooks(built, expected):
    """Compare two workbooks sheet by sheet and cell by cell (values, exact equality).

    Parameters
    ----------
    built, expected : pathlib.Path
        Workbooks to compare.

    Returns
    -------
    dict
        ``status`` (``IDENTICAL`` or ``DIFFERENT``), sheet names, cell counts and up to 20
        differing cells.
    """
    import openpyxl
    a = openpyxl.load_workbook(built); b = openpyxl.load_workbook(expected)
    diffs, n = [], 0
    if a.sheetnames != b.sheetnames:
        diffs.append(('sheetnames', a.sheetnames, b.sheetnames))
    for name in b.sheetnames:
        if name not in a.sheetnames:
            continue
        ra = list(a[name].iter_rows(values_only=True)); rb = list(b[name].iter_rows(values_only=True))
        if len(ra) != len(rb):
            diffs.append((name, 'rows', len(ra), len(rb)))
        for i, (x, y) in enumerate(zip(ra, rb)):
            if len(x) != len(y):
                diffs.append((name, i + 1, 'columns', len(x), len(y)))
            for j, (u, v) in enumerate(zip(x, y)):
                n += 1
                if u != v or type(u) is not type(v):
                    diffs.append((name, i + 1, j + 1, repr(u), repr(v)))
    return dict(status='IDENTICAL' if not diffs else 'DIFFERENT', sheets=b.sheetnames,
                cells_compared=n, n_differences=len(diffs), differences=[list(map(str, d)) for d in diffs[:20]])


def read_endpoints(sd):
    """Read the Endpoints table from source_data (published values)."""
    return pd.read_csv(sd / 'endpoints_by_definition.tsv', sep='\t')


# ---------------------------------------------------------------------------------
# Controlled route: the statistics
# ---------------------------------------------------------------------------------
def controlled_path(name):
    """Return the path of a controlled input, or None (printing why) when it is absent.

    Parameters
    ----------
    name : str
        File name listed in controlled/CONTROLLED_INPUTS.md.
    """
    root = os.environ.get('FASTALAKE_CONTROLLED_DIR')
    p = Path(root) / name if root else None
    if p is None or not p.is_file():
        print(f'SKIPPED: needs controlled input {name} (see controlled/CONTROLLED_INPUTS.md)')
        return None
    return p


def bh(p):
    """Benjamini-Hochberg q values (moved unchanged from the producer)."""
    p = np.asarray(p, float); o = np.argsort(p); r = p[o]*len(p)/np.arange(1, len(p)+1)
    q = np.empty_like(p); q[o] = np.minimum(1, np.minimum.accumulate(r[::-1])[::-1]); return q


def derive_state(d, extra_core=()):
    """Antibiotic state per acquisition (On, Off, Unknown), moved unchanged.

    On if any core flag or any ``visit_antibiotic_treatment_*`` flag is '1'; Off if all
    core flags are '0'; otherwise Unknown.
    """
    core = list(CORE_FLAGS) + list(extra_core)
    flags = core + [c for c in d if c.startswith('visit_antibiotic_treatment') and c != 'visit_antibiotic_treatment_other_specify']
    assert set(d[flags].stack()) <= {'', '0', '1'}
    pos = d[flags].eq('1').any(axis=1)
    return pd.Series(np.where(pos, 'On', np.where(d[core].eq('0').all(axis=1), 'Off', 'Unknown')), index=d.index)


def paired_inference(ac, state):
    """Paired on-minus-off inference for one exposure definition (moved unchanged).

    Statistical unit: patient. n: patients with acquisitions in both known states
    (42, 40, 35 and 38 for definitions A to D). Acquisitions are combined within patient
    and state (composition fractions averaged, protein-group counts by median).

    Parameters
    ----------
    ac : pandas.DataFrame
        Acquisition table (controlled ``acquisition_accounting.tsv``), 306 rows.
    state : pandas.Series
        State per acquisition, aligned with ``ac``.

    Returns
    -------
    tuple
        (statistics per endpoint, composition by state, paired patient table, acquisitions by state)
    """
    a = ac.assign(antibiotic_state=state.to_numpy())
    st = a[a.antibiotic_state.ne('Unknown')].groupby(['patient_id', 'antibiotic_state'])
    paired = st[CLASSES].mean().join(st.pg_parsimony.median()).join(st.size().rename('n_acquisitions')).reset_index()
    el = paired.groupby('patient_id').antibiotic_state.nunique()
    paired = paired[paired.patient_id.isin(el[el.eq(2)].index)]
    n = paired.patient_id.nunique(); assert len(paired) == 2*n
    rows = []
    for idx, (metric, scale) in enumerate(ENDPOINTS):
        wide = paired.pivot(index='patient_id', columns='antibiotic_state', values=metric)[['Off', 'On']]*scale
        delta = (wide.On-wide.Off).to_numpy()
        rng = np.random.default_rng(SEED)
        boots = delta[rng.integers(0, len(delta), (N_BOOTSTRAP, len(delta)))].mean(axis=1)
        lo, hi = np.quantile(boots, list(CI_QUANTILES))
        rng = np.random.default_rng(SEED+idx); count = 0
        for _ in range(N_SIGNFLIP_BATCHES):
            null = ((rng.integers(0, 2, (SIGNFLIP_BATCH, len(delta)))*2-1)*delta).mean(axis=1)
            count += int((abs(null) >= abs(delta.mean())-TIE_TOLERANCE).sum())
        total = N_SIGNFLIP_BATCHES*SIGNFLIP_BATCH
        rows.append(dict(metric=metric, n_patients=n, mean_off=wide.Off.mean(), mean_on=wide.On.mean(),
                         mean_on_minus_off=delta.mean(), median_on_minus_off=float(np.median(delta)),
                         n_patients_increase=int((delta > 0).sum()), CI95_low=lo, CI95_high=hi,
                         p_signflip_MC=(count+PSEUDOCOUNT)/(total+PSEUDOCOUNT),
                         signflip_extreme_count=count, signflip_draws=total))
    s = pd.DataFrame(rows); s['q_two_clinical_endpoints'] = bh(s.p_signflip_MC)
    comp = paired.groupby('antibiotic_state')[CLASSES].mean().reset_index()
    acq = paired.groupby('antibiotic_state').n_acquisitions.sum().to_dict()
    return s, comp, paired, acq


def load_controlled():
    """Read and hash-check the two controlled inputs.

    Returns
    -------
    tuple or None
        ``(d, ac)``: metadata rows in acquisition order and the acquisition table, or
        None (with a printed reason) when a controlled input is absent.
    """
    paths = {k: controlled_path(k) for k in CONTROLLED}
    if any(p is None for p in paths.values()):
        return None
    for k, p in paths.items():
        assert sha(p) == CONTROLLED[k], f'controlled input {k} is not the file the published values came from'
    d = pd.read_csv(paths['PREDICT_metadata_2026-08-03.tsv'], sep='\t', dtype=str, keep_default_na=False).set_index('sample_id')
    ac = pd.read_csv(paths['acquisition_accounting.tsv'], sep='\t', dtype={'patient_id': str})
    assert len(ac) == ac.sample_id.nunique() == N_ACQUISITIONS
    return d.loc[ac.sample_id], ac


def stat_definitions(d, ac, primary, negative_control=False):
    """All Supplementary Table 3 statistics and counts (port of ``rifaximin_sensitivity.py::main``).

    Statistical unit: patient; n = 42, 40, 35, 38 paired patients for definitions A-D,
    from 306 acquisitions. Per definition: two endpoints (human-assigned share,
    percentage points; quantified protein groups), mean paired on-minus-off difference,
    50,000-draw percentile patient bootstrap (seed 20260907), 200,000 two-sided sign flips
    (seed 20260907 + endpoint index) with plus-one correction, Benjamini-Hochberg across
    the two endpoints.

    Parameters
    ----------
    d : pandas.DataFrame
        Controlled metadata rows aligned with ``ac``.
    ac : pandas.DataFrame
        Controlled acquisition table.
    primary : pandas.DataFrame
        ``source_data/STab3/primary_paired_statistics.tsv``; scenario A must reproduce it.
    negative_control : bool
        Invert every HE flag (planted control; results are not for reporting).

    Returns
    -------
    dict
        ``stats`` (DataFrame, 8 rows), ``composition`` (DataFrame, 8 rows), ``counts`` (dict).
    """
    he = d[HE_FIELD]
    if negative_control:
        he = he.map({'0': '1', '1': '0', '': ''})
        d = d.assign(**{HE_FIELD: he})
    sA = derive_state(d)
    gate = (sA.to_numpy() == ac.antibiotic_state.to_numpy())
    assert gate.all(), f'{(~gate).sum()} state mismatches against the acquisition table'
    sB = derive_state(d, extra_core=[HE_FIELD])
    he1 = he.eq('1').to_numpy()
    pid = ac.patient_id.to_numpy()
    he_pat = set(pid[he1])
    sC = pd.Series(np.where(np.isin(pid, list(he_pat)), 'Unknown', sA), index=d.index)
    sD = pd.Series(np.where(he1 & (sA.to_numpy() == 'Off'), 'Unknown', sA), index=d.index)
    scen = dict(zip(SCENARIOS, [sA, sB, sC, sD]))

    counts = dict(
        acquisitions=N_ACQUISITIONS,
        HE_flag_values={k: int(v) for k, v in he.value_counts().items()},
        HE_by_current_state={f'{h}|{s}': int(v) for (h, s), v in pd.crosstab(he, sA).stack().items()},
        HE_rows_patients=len(he_pat),
        HE_off_rows=int((he1 & (sA.to_numpy() == 'Off')).sum()),
        HE_off_row_patients=int(len(set(pid[he1 & (sA.to_numpy() == 'Off')]))),
        state_changes_A_to_B={f'{a}->{b}': int(v) for (a, b), v in pd.crosstab(sA, sB).stack().items() if v and a != b},
        HE_with_core_rifaximin=int((he1 & d.visit_rifaximin.eq('1').to_numpy()).sum()),
        HE_with_treatment_rifaximin=int((he1 & d.visit_antibiotic_treatment_rifaximin.eq('1').to_numpy()).sum()),
    )
    allstats, allcomp, membership = [], [], {}
    for name, s in scen.items():
        st, comp, paired, acq = paired_inference(ac, s)
        st.insert(0, 'scenario', name); comp.insert(0, 'scenario', name)
        st['acq_off'] = acq.get('Off'); st['acq_on'] = acq.get('On')
        allstats.append(st); allcomp.append(comp); membership[name] = set(paired.patient_id)
    A = membership['A_current']
    for name, m in membership.items():
        counts[f'paired_{name}'] = dict(n=len(m), dropped_vs_A=len(A-m), added_vs_A=len(m-A),
                                        HE_patients_in_pairs=len(m & he_pat))
    counts['HE_patients_in_A_pairs'] = len(A & he_pat)
    counts['A_pairs_with_HE_off_rows'] = len(A & set(pid[he1 & (sA.to_numpy() == 'Off')]))
    stats = pd.concat(allstats, ignore_index=True); comp = pd.concat(allcomp, ignore_index=True)

    a = stats[stats.scenario.eq('A_current')].reset_index(drop=True)
    for c in ['n_patients', 'mean_off', 'mean_on', 'mean_on_minus_off', 'CI95_low', 'CI95_high', 'p_signflip_MC', 'q_two_clinical_endpoints']:
        assert np.allclose(a[c], primary[c], rtol=1e-10, atol=1e-12), c
    counts['gate_A_reproduces_clinical_refit'] = True
    return dict(stats=stats, composition=comp, counts=counts)


def flatten_counts(counts, prefix=''):
    """Flatten nested counts into ``{key: value}`` with '.'-joined keys (as in definition_counts.tsv)."""
    out = {}
    for k, v in counts.items():
        if isinstance(v, dict):
            out.update(flatten_counts(v, f'{prefix}{k}.'))
        else:
            out[f'{prefix}{k}'] = v
    return out


def compare_to_source(res, sd):
    """Compare recomputed statistics, composition and counts with source_data.

    Returns
    -------
    dict
        Maximum relative difference per table and the count comparison.
    """
    prod = pd.read_csv(sd / 'endpoints_producer_output.tsv', sep='\t', float_precision='round_trip')
    pub = read_endpoints(sd)
    s = res['stats']
    num = ['n_patients', 'mean_off', 'mean_on', 'mean_on_minus_off', 'median_on_minus_off', 'n_patients_increase',
           'CI95_low', 'CI95_high', 'p_signflip_MC', 'q_two_clinical_endpoints', 'acq_off', 'acq_on']
    exact_prod = bool((s[['scenario', 'metric']].values == prod[['scenario', 'metric']].values).all()
                      and all(np.array_equal(s[c].to_numpy(float), prod[c].to_numpy(float)) for c in num))
    pubmap = dict(n_patients_increase='n_patients_on_greater_than_off', q_two_clinical_endpoints='q_BH_two_endpoints_within_definition',
                  acq_off='n_acquisitions_off', acq_on='n_acquisitions_on')
    rel_pub = max(float(np.max(np.abs(s[c].to_numpy(float) - pub[pubmap.get(c, c)].to_numpy(float))
                               / np.maximum(np.abs(pub[pubmap.get(c, c)].to_numpy(float)), 1e-300))) for c in num)
    comp = pd.read_csv(sd / 'composition_by_definition.tsv', sep='\t', float_precision='round_trip')
    rc = res['composition']
    exact_comp = bool((rc.scenario.values == comp.scenario.values).all() and (rc.antibiotic_state.values == comp.state.values).all()
                      and all(np.array_equal(rc[c].to_numpy(float), comp[c].to_numpy(float)) for c in CLASSES))
    cnt = pd.read_csv(sd / 'definition_counts.tsv', sep='\t', dtype=str, keep_default_na=False)
    flat = {k: str(v) for k, v in flatten_counts(res['counts']).items()}
    counts_equal = flat == dict(zip(cnt['count'], cnt['value']))
    return dict(stats_identical_to_producer_output=exact_prod, stats_max_rel_diff_vs_published_table=rel_pub,
                composition_identical=exact_comp, counts_identical=counts_equal)


def numbers(res):
    """Return the printed numbers of this unit in NUMBERS.tsv order, from recomputed results.

    Parameters
    ----------
    res : dict
        Output of :func:`stat_definitions`.

    Returns
    -------
    list of tuple
        ``(label, value)`` with full precision.
    """
    s = res['stats'].set_index(['scenario', 'metric']); c = res['counts']
    B = 'B_HE_as_exposure'
    n = {k: int(s.loc[(k, 'Human'), 'n_patients']) for k in SCENARIOS}
    rows = [
        ('137 B paired patients', n[B]),
        ('137 B human share mean on-off (pp)', s.loc[(B, 'Human'), 'mean_on_minus_off']),
        ('137 B human share CI low', s.loc[(B, 'Human'), 'CI95_low']),
        ('137 B human share CI high', s.loc[(B, 'Human'), 'CI95_high']),
        ('137 B human share q', s.loc[(B, 'Human'), 'q_two_clinical_endpoints']),
        ('137 B depth fewer groups (-mean on-off)', -s.loc[(B, 'pg_parsimony'), 'mean_on_minus_off']),
        ('137 B depth CI, smaller reduction (-CI high)', -s.loc[(B, 'pg_parsimony'), 'CI95_high']),
        ('137 B depth CI, larger reduction (-CI low)', -s.loc[(B, 'pg_parsimony'), 'CI95_low']),
        ('137 B depth q', s.loc[(B, 'pg_parsimony'), 'q_two_clinical_endpoints']),
        ('479 A patients', n['A_current']),
        ('479 B patients', n[B]),
        ('479 C patients', n['C_HE_patients_excluded']),
        ('479 D patients', n['D_HE_off_rows_excluded']),
        ('793 Off acquisitions reclassified A->B', c['state_changes_A_to_B'].get('Off->On', 0)),
        ('793 patients of those acquisitions', NUMBER_WORDS.get(c['HE_off_row_patients'], c['HE_off_row_patients'])),
        ('793 B paired patients', n[B]),
        ('793 C patients', n['C_HE_patients_excluded']),
        ('793 A patients', n['A_current']),
        ('793 Off acquisitions with HE record set Unknown (D)', c['HE_off_rows']),
        ('793 D patients', n['D_HE_off_rows_excluded']),
        ('893 endpoints per BH family', NUMBER_WORDS.get(len(ENDPOINTS), len(ENDPOINTS))),
    ]
    return [(k, v if isinstance(v, str) else int(v) if isinstance(v, (int, np.integer)) else float(v))
            for k, v in rows]
