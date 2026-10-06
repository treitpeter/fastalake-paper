"""Supplementary Table 5: individual-function correlations between DNA, RNA and protein layers.

PREDICT cohort, 77 patients. For each functional representation (OG, KO, pathway) and
each pair of microbial layers (metagenome DNA, metatranscriptome RNA, non-host protein),
one Spearman correlation across patients per function, Benjamini-Hochberg within each of
the nine representation-by-layer-pair families (Supplementary Methods 8, v28d line 829).

This module does three things, each moved unchanged from the original code (paths made
repository-relative, script bodies split into functions):

1. ``stat_individual_functions`` recomputes all 30,581 tests, the eligibility table and the
   nine-family summary from the patient-level profiles. Origin:
   ``LAB_ROOT/.../biology_refit/individual_functions.py`` (refit mode).
2. ``stat_eligibility_common`` matches revised and previous-allocation tests by
   representation, layer pair and function. Origin: Supplementary Table 4 block of
   ``LAB_ROOT/.../iteration2/cr003_followup/scripts/tables_cr003.py``.
3. ``build_table`` writes the workbook as the producer did (same script) and applies the
   Source Data package edit (Read_me sheet, column renames) of
   ``LAB_ROOT/.../final/figure_sd/source_data/build/specials.py`` (functions ``st4``,
   ``elig``, ``elig_sum``).

Source data (``source_data/STab5/``):

- ``profiles_<level>_<layer>.tsv.gz``  patient-level inputs of the correlation step
- ``crossomics_same_function_correlations.tsv``  the 30,581 tests as produced
- ``same_function_summary.tsv``, ``eligibility.tsv``  as produced
- ``previous_allocation_same_function_correlations.tsv``  the 17,914 tests of the
  previous allocation (frozen input of the eligibility comparison)
"""

import io
from itertools import combinations
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import spearmanr
from statsmodels.stats.multitest import multipletests

# ---------------------------------------------------------------------------------------
# Choices (settings of the upstream analysis)
# ---------------------------------------------------------------------------------------
#: Patients with complete DNA, RNA, protein and metabolite data (Supplementary Methods 8, v28d line 829;
#: one earliest complete specimen per patient, Supplementary Methods 8, v28d line 809).
N_PATIENTS = 77
#: Minimum patients with positive values in both layers (Supplementary Methods 8, v28d line 829).
MIN_PATIENTS = 35
#: Functional representations (Supplementary Methods 8, v28d line 829; definitions Supplementary Methods 5, v28d line 657).
LEVELS = ["OG", "KO", "Pathway"]
#: Molecular layers; every pair is tested (Supplementary Methods 8, v28d line 829).
LAYERS = ["metaG", "metaT", "microbial_metaP"]
#: Correlation (Supplementary Methods 8, v28d line 829); two-sided P from scipy.stats.spearmanr.
CORRELATION = "spearman"
#: Multiple-testing correction, one family per level x layer pair (Supplementary Methods 8, v28d line 829;
#: Supplementary Methods 10, v28d line 893).
CORRECTION = "fdr_bh"
FAMILY_KEYS = ["level", "layer_a", "layer_b"]
#: Significance threshold for the Nine_family_summary column `significant` (producer).
Q_SIGNIFICANT = 0.05
#: Keys that match revised and previous-allocation tests (Supplementary Methods 8, v28d line 829).
MATCH_KEYS = ["level", "layer_a", "layer_b", "feature"]
#: Relative tolerance for float comparisons with the tables as produced (producer value).
FLOAT_RTOL = 1e-12

# Producer note sheet (tables_cr003.py, non-replay variant) and package edits (specials.py).
PRODUCER_README = [
    'Individual-function Spearman correlations between microbial DNA, RNA and protein layers across PREDICT patients ({ntests:,} tests).',
    'Native UHGG eggNOG annotations were used only after the protein aligned to the annotation’s seed orthologue with at least half of the reported bitscore.',
    'Benjamini–Hochberg correction was applied separately within each of the nine representation-by-layer-pair families.',
    'Eligibility_common compares functions eligible under the previous and revised allocations, matched by representation, layer pair and function identifier.',
    'No profiles were renormalized and no correlations or p/q values were refitted for the common-eligibility subset.',
]
PACKAGE_README = [
    'Supplementary Table 5. Individual-function correlations between DNA, RNA and protein layers in the PREDICT cohort.',
    '',
    'All_30581_tests: one row per eligible function and layer pair under the revised allocation (30,581 tests): number of patients, Spearman rho, two-sided P value and Benjamini-Hochberg q value within each of the nine representation-by-layer-pair families.',
    'Nine_family_summary: for each family, the number of tested functions, the median rho and the number of functions with q < 0.05.',
    'All_eligibility: every function and layer pair considered, with the number of patients with positive measurements in both layers and whether it met the eligibility rule (at least 35 patients).',
    'Eligibility_common: for each family, the number of functions eligible under both the previous and the revised allocation (n_common), eligible only under the revised allocation (n_newly_eligible) and eligible only under the previous allocation (n_lost), with the median revised correlation for common and for newly eligible functions and the median previous correlation for common functions. Profiles were not renormalised, and correlations, P values and q values were not refitted for this subset.',
    '',
    'Previous allocation: the protein-to-function allocation used before the annotation revision of 2 October 2026. Revised allocation: the allocation used for all other values in this table.',
    'Native UHGG eggNOG annotations were used only when the protein aligned to the annotation seed orthologue with at least half of the reported bitscore (Methods). Values computed before this check are in Source Data, Annotation_verification.',
    'Per-function rows of the eligibility comparison: Source Data, Supplementary_Tables/Table5_support/eligibility_matched_correlations.tsv.',
]
PACKAGE_SHEETS = ('All_30581_tests', 'Nine_family_summary', 'All_eligibility', 'Eligibility_common')
ELIG_COLS = {'n_patients_refit': 'n_patients_revised_allocation', 'rho_refit': 'rho_revised_allocation',
             'n_patients_baseline': 'n_patients_previous_allocation', 'rho_baseline': 'rho_previous_allocation'}
ELIG_VALS = {'both': 'common', 'left_only': 'newly_eligible', 'right_only': 'lost'}
SUM_COLS = {'n_new': 'n_newly_eligible', 'median_refit_common': 'median_rho_revised_common',
            'median_baseline_common': 'median_rho_previous_common', 'median_refit_new': 'median_rho_revised_newly_eligible'}


def rd(p, **kw):
    """Read a TSV as the original code did (pandas default float parser)."""
    return pd.read_csv(p, sep="\t", **kw)


def load(sd):
    """Read every table of ``source_data/STab5/``.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab5``.

    Returns
    -------
    dict
        ``tests``, ``summary``, ``eligibility``, ``previous`` (DataFrames as produced).
    """
    return dict(tests=rd(sd / "crossomics_same_function_correlations.tsv"),
                summary=rd(sd / "same_function_summary.tsv"),
                eligibility=rd(sd / "eligibility.tsv"),
                previous=rd(sd / "previous_allocation_same_function_correlations.tsv"))


def load_profiles(sd, level):
    """Rebuild the patient x feature matrices of one level from the long profile files.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab5``.
    level : str
        ``OG``, ``KO`` or ``Pathway``.

    Returns
    -------
    dict of str to pandas.DataFrame
        One dense matrix per layer (77 pseudonymised patients, sorted, x every feature in
        ``profile_features.tsv``; values absent from the long file are zero).
    """
    patients = pd.read_csv(sd / "patients.tsv", sep="\t").patient.tolist()
    assert len(patients) == len(set(patients)) == N_PATIENTS
    vocab = pd.read_csv(sd / "profile_features.tsv", sep="\t", keep_default_na=False)
    out = {}
    for lay in LAYERS:
        d = pd.read_csv(sd / f"profiles_{level}_{lay}.tsv.gz", sep="\t", float_precision="round_trip",
                        keep_default_na=False)
        cols = vocab[(vocab.level == level) & (vocab.layer == lay)].feature.tolist()
        assert set(d.feature) <= set(cols) and set(d.patient) <= set(patients)
        m = d.pivot(index="patient", columns="feature", values="value")
        out[lay] = m.reindex(index=patients, columns=cols).fillna(0.0)
    return out


def relative(d):
    """Original ``relative``: divide each row by its sum over the given columns."""
    return d.div(d.sum(1).replace(0, np.nan), axis=0)


def stat_individual_functions(sd):
    """Recompute every individual-function Spearman test, the eligibility table and the summary.

    Statistical unit: patient (n = 77 PREDICT patients; each test uses the patients with
    positive values in both layers, at least ``MIN_PATIENTS``). Spearman correlation,
    two-sided P (scipy.stats.spearmanr), Benjamini-Hochberg within each of the nine
    level x layer-pair families (statsmodels ``multipletests``, method ``fdr_bh``).

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab5``.

    Returns
    -------
    tuple of pandas.DataFrame
        ``(tests, eligibility, summary)`` with the producer's columns and row order.
    """
    rows, elig = [], []
    for lev in LEVELS:
        fs = load_profiles(sd, lev)
        for a, b in combinations(fs, 2):
            terms = sorted(set(fs[a]) & set(fs[b]))
            aa, bb = [relative(fs[k].loc[:, terms]) for k in [a, b]]
            for term in terms:
                ok = (aa[term] > 0) & (bb[term] > 0)
                n = int(ok.sum())
                r = p = np.nan
                if n >= MIN_PATIENTS:
                    r, p = spearmanr(aa.loc[ok, term], bb.loc[ok, term])
                elig.append(dict(level=lev, layer_a=a, layer_b=b, feature=term, n_patients=n, eligible=bool(np.isfinite(p))))
                if np.isfinite(p):
                    rows.append(dict(level=lev, layer_a=a, layer_b=b, feature=term, n_patients=n, rho=r, p=p))
    res = pd.DataFrame(rows)
    res["q"] = res.groupby(FAMILY_KEYS).p.transform(lambda v: multipletests(v, method=CORRECTION)[1])
    summary = summarise(res)
    return res, pd.DataFrame(elig), summary


def summarise(res):
    """Nine-family summary (original code): tests, median rho and tests with q < 0.05.

    Parameters
    ----------
    res : pandas.DataFrame
        Test table (``level``, ``layer_a``, ``layer_b``, ``rho``, ``q``).

    Returns
    -------
    pandas.DataFrame
        One row per family.
    """
    return res.groupby(FAMILY_KEYS).agg(features=('rho', 'size'), median_rho=('rho', 'median'),
                                        significant=('q', lambda v: int((v < Q_SIGNIFICANT).sum()))).reset_index()


def stat_eligibility_common(ref, base):
    """Match revised and previous-allocation tests and summarise each family (original code).

    Statistical unit: function x layer pair (descriptive medians; no test). Revised tests:
    30,581; previous-allocation tests: 17,914; both from 77 patients.

    Parameters
    ----------
    ref : pandas.DataFrame
        Revised-allocation tests.
    base : pandas.DataFrame
        Previous-allocation tests.

    Returns
    -------
    tuple of pandas.DataFrame
        ``(matched, ver)``: the per-function outer join with indicator ``eligibility`` and
        the per-family counts and medians (producer column names).
    """
    m = ref.merge(base[MATCH_KEYS + ['n_patients', 'rho']], on=MATCH_KEYS, how='outer', suffixes=('_refit', '_baseline'),
                  indicator='eligibility', validate='one_to_one')
    ver = []
    for (l, a, b), g in m.groupby(FAMILY_KEYS, sort=True):
        both = g[g.eligibility == 'both']
        new = g[g.eligibility == 'left_only']
        ver.append(dict(level=l, layer_a=a, layer_b=b, n_common=len(both), n_new=len(new),
                        n_lost=int((g.eligibility == 'right_only').sum()),
                        median_refit_common=both.rho_refit.median(), median_baseline_common=both.rho_baseline.median(),
                        median_refit_new=new.rho_refit.median() if len(new) else np.nan))
    return m, pd.DataFrame(ver)


def package_support(m, ver):
    """Apply the Source Data package edits to the two Table 5 support files (specials.py).

    Parameters
    ----------
    m, ver : pandas.DataFrame
        Outputs of ``stat_eligibility_common``.

    Returns
    -------
    tuple of bytes
        ``(eligibility_matched_correlations.tsv, common_eligibility_summary.tsv)`` as in
        the Source Data package.
    """
    import csv
    def tsv_bytes(df):
        return df.to_csv(sep='\t', index=False).encode('utf-8')
    def _tsv(data):
        return list(csv.reader(io.StringIO(data.decode('utf-8')), delimiter='\t'))
    def _write(rows):
        out = io.StringIO(); csv.writer(out, delimiter='\t', lineterminator='\n').writerows(rows)
        return out.getvalue().encode('utf-8')
    rows = _tsv(tsv_bytes(m))
    h = rows[0]
    rows[0] = [ELIG_COLS.get(x, x) for x in h]
    j = h.index('eligibility')
    for r in rows[1:]:
        r[j] = ELIG_VALS[r[j]]
    elig = _write(rows)
    rows = _tsv(tsv_bytes(ver))
    rows[0] = [SUM_COLS.get(x, x) for x in rows[0]]
    return elig, _write(rows)


def build_table(sd, out):
    """Write Supplementary Table 5 as the producer did, then apply the package edit.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab5``.
    out : pathlib.Path
        Output directory.

    Returns
    -------
    tuple
        ``(xlsx_path, matched, ver)``.
    """
    import openpyxl
    out.mkdir(parents=True, exist_ok=True)
    ref = rd(sd / 'crossomics_same_function_correlations.tsv')
    base = rd(sd / 'previous_allocation_same_function_correlations.tsv')
    m, ver = stat_eligibility_common(ref, base)
    ntests = len(ref)
    producer = out / 'producer_Supplementary_Table_5.xlsx'
    with pd.ExcelWriter(producer, engine='openpyxl') as w:
        for file, sheet in [('crossomics_same_function_correlations.tsv', f'All_{ntests}_tests'),
                            ('same_function_summary.tsv', 'Nine_family_summary'), ('eligibility.tsv', 'All_eligibility')]:
            rd(sd / file).to_excel(w, sheet_name=sheet, index=False)
    book = openpyxl.load_workbook(producer)
    notes = book.create_sheet('Read_me', 0)
    for line in PRODUCER_README:
        notes.append([line.format(ntests=ntests)])
    sh = book.create_sheet('Eligibility_common'); sh.append(list(ver.columns)); sh.freeze_panes = 'A2'
    for row in ver.itertuples(index=False, name=None):
        sh.append([None if pd.isna(v) else v for v in row])
    book.save(producer)
    # Source Data package edit (specials.st4), then neutral document properties.
    wb = openpyxl.load_workbook(producer)
    for n in [n for n in wb.sheetnames if n not in PACKAGE_SHEETS]:
        del wb[n]
    ws = wb.create_sheet('Read_me', 0)
    for line in PACKAGE_README:
        ws.append([line])
    for c in wb['Eligibility_common'][1]:
        if c.value in SUM_COLS:
            c.value = SUM_COLS[c.value]
    wb.properties.creator = 'FastaLake'; wb.properties.lastModifiedBy = 'FastaLake'
    xlsx = out / 'Supplementary_Table_5.xlsx'
    wb.save(xlsx)
    return xlsx, m, ver


def sheets(p):
    """All cell values of a workbook, by sheet."""
    import openpyxl
    b = openpyxl.load_workbook(p, data_only=True)
    return {s.title: [tuple(r) for r in s.values] for s in b}


def compare_xlsx_values(new, expected):
    """Compare two workbooks cell by cell (sheet names, order, shape, values).

    Floats agree if equal or within ``FLOAT_RTOL``; the receipt counts both kinds.

    Returns
    -------
    dict
        ``identical`` (bool), ``sheets``, ``cells``, ``exact_cells``, ``first_difference``.
    """
    a, b = sheets(new), sheets(expected)
    res = dict(sheets=list(a), expected_sheets=list(b), cells=0, exact_cells=0, first_difference=None)
    if list(a) != list(b):
        res['first_difference'] = 'sheet names or order'
    else:
        for k in a:
            if len(a[k]) != len(b[k]):
                res['first_difference'] = f'{k}: rows {len(a[k])} vs {len(b[k])}'; break
            for i, (r1, r2) in enumerate(zip(a[k], b[k])):
                if len(r1) != len(r2):
                    res['first_difference'] = f'{k} row {i}: width'; break
                for x, y in zip(r1, r2):
                    res['cells'] += 1
                    if x == y:
                        res['exact_cells'] += 1; continue
                    if isinstance(x, float) and isinstance(y, float) and (np.isclose(x, y, rtol=FLOAT_RTOL, atol=0) or (np.isnan(x) and np.isnan(y))):
                        continue
                    res['first_difference'] = f'{k} row {i}: {x!r} vs {y!r}'; break
                if res['first_difference']:
                    break
            if res['first_difference']:
                break
    res['identical'] = res['first_difference'] is None
    return res


def compare_frames(new, old, keys, floats, ints=()):
    """Compare a recomputed table with the table as produced, matched on ``keys``.

    Returns
    -------
    dict
        Row counts, key agreement, max absolute float difference per column, integer and
        boolean agreement.
    """
    j = new.merge(old, on=list(keys), how='outer', suffixes=('_new', '_old'), indicator=True)
    out = dict(rows_new=len(new), rows_old=len(old), keys_equal=bool(j._merge.eq('both').all()))
    for c in floats:
        d = (j[c + '_new'] - j[c + '_old']).abs()
        out[f'max_abs_diff_{c}'] = float(d.max())
    for c in ints:
        out[f'equal_{c}'] = bool((j[c + '_new'] == j[c + '_old']).all())
    return out


def numbers(summary, ver, tests, previous):
    """Every printed value this unit covers, in NUMBERS.tsv order (full precision).

    Statistical unit: function x layer pair; medians across functions of per-function
    Spearman rho (77 patients).

    Returns
    -------
    list of tuple
        ``(label, value)``.
    """
    prot = summary[summary.layer_b == 'microbial_metaP']
    dr = summary[(summary.layer_a == 'metaG') & (summary.layer_b == 'metaT')]
    vp = ver[ver.layer_b == 'microbial_metaP']
    words = {9: 'nine'}
    return [
        ('L141 DNA/RNA-protein median rho, min', float(prot.median_rho.min())),
        ('L141 DNA/RNA-protein median rho, max', float(prot.median_rho.max())),
        ('L141 DNA-RNA median rho, min', float(dr.median_rho.min())),
        ('L141 DNA-RNA median rho, max', float(dr.median_rho.max())),
        ('L141 common-eligible protein median rho (revised), min', float(vp.median_refit_common.min())),
        ('L141 common-eligible protein median rho (revised), max', float(vp.median_refit_common.max())),
        ('L829 smallest number of patients in any test (rule: at least MIN_PATIENTS)', int(tests.n_patients.min())),
        ('L829 BH families', words[int(tests.groupby(FAMILY_KEYS).ngroups)]),
        ('L829 tests under the revised allocation', int(len(tests))),
        ('L829 common-eligible protein median rho (previous), min', float(vp.median_baseline_common.min())),
        ('L829 common-eligible protein median rho (previous), max', float(vp.median_baseline_common.max())),
        ('L893 BH families for individual-function correlations', words[int(tests.groupby(FAMILY_KEYS).ngroups)]),
    ]


def stat_families(tests):
    """Per-family description of the Benjamini-Hochberg families (for STATS.tsv).

    Statistical unit: patient within each test (n = 35-77 patients with positive values
    in both layers, of 77); one row per level x layer-pair family.

    Parameters
    ----------
    tests : pandas.DataFrame
        The 30,581-row test table.

    Returns
    -------
    pandas.DataFrame
        ``family_size``, minimum and maximum patients per test, median rho, interquartile
        range of rho, number with q < 0.05, minimum P and q.
    """
    g = tests.groupby(FAMILY_KEYS)
    return pd.DataFrame(dict(
        family_size=g.size(), n_min=g.n_patients.min(), n_max=g.n_patients.max(),
        median_rho=g.rho.median(), rho_q25=g.rho.quantile(0.25), rho_q75=g.rho.quantile(0.75),
        significant=g.q.apply(lambda v: int((v < Q_SIGNIFICANT).sum())), min_p=g.p.min(), min_q=g.q.min(),
    )).reset_index()
