"""Supplementary Table 1: candidate-clustering search sensitivity in 15 CAMPI acquisitions.

Builds the 105-row table (15 acquisitions x 7 arms: unclustered and six MMseqs2/Linclust
identity thresholds), repeats the producer's consistency checks and computes the medians
printed in the Results (v28d line 89).

The table logic is moved unchanged from the original producer
``LAB_ROOT/.../final/supp_table1/build_supp_table1.py`` (see PROVENANCE.md). Only the
input and output paths were made repository-relative and the script body was split into
functions.

Source data (``source_data/STab1/``):

- ``per_acquisition_identifications.tsv``  producer table, 105 rows (peptide counts)
- ``per_acquisition_KO_coverage.tsv``      producer table, 105 rows (KO counts)
- ``support_97_99_identification_comparison.tsv``  the 45 rows (unclustered, 97%, 99%)
  of the earlier Supplementary Table 1, used only as a check
- ``check_new_rows_95_90_75_50.tsv``       the 60 rows added for 95/90/75/50%, used only
  as a check
- ``independent_recount_identifications.tsv``, ``independent_recount_KO.tsv``
  an independent recount from the Sage result files, used only as a check
"""

from pathlib import Path

import pandas as pd

# ---------------------------------------------------------------------------------------
# Choices (each is a setting of the upstream analysis, not of this table builder)
# ---------------------------------------------------------------------------------------
#: Arms and their Linclust --min-seq-id in percent (Methods, v28d line 695).
ARM_ID = {"unclustered": None, "cluster_099": 99, "cluster_097": 97, "cluster_095": 95,
          "cluster_090": 90, "cluster_075": 75, "cluster_050": 50}
#: Linclust coverage (-c) and coverage mode (Methods, v28d line 695); recorded only.
LINCLUST_COVERAGE = 0.8
LINCLUST_COV_MODE = 0
LINCLUST_CLUSTER_MODE = 2
#: Peptide acceptance (Methods, v28d line 695; legend v28d line 475 via LEGEND.txt).
PEPTIDE_Q_CUTOFF = 0.01
PEPTIDE_LENGTH = (5, 50)
#: Number of CAMPI acquisitions (legend, v28d line 475).
N_ACQUISITIONS = 15
#: Expected table rows: acquisitions x arms (Methods, v28d line 695).
N_ROWS = 105
#: Tolerance for fraction columns in the cross-checks (producer value).
FLOAT_TOL = 1e-12

ARM_ORDER = {a: i for i, a in enumerate(ARM_ID)}

#: Producer column -> column of the earlier 45-row table (producer's MAP).
MAP = {
    "peptides": "accepted", "REF_peptides": "reference_denominator", "REF_recovered": "reference_recovered",
    "lost": "lost_vs_unclustered", "gained": "gained_vs_unclustered", "peptide_count_ratio": "count_ratio",
    "REF_recovery": "reference_recovery_fraction", "baseline_identity_recovery": "baseline_identity_retained"}
INTS = ["accepted", "reference_denominator", "reference_recovered", "lost_vs_unclustered", "gained_vs_unclustered"]
FLOATS = ["count_ratio", "reference_recovery_fraction", "baseline_identity_retained"]


def _read(path):
    return pd.read_csv(path, sep="\t", float_precision="round_trip")


def load(sd):
    """Read every source table this unit uses.

    Parameters
    ----------
    sd : pathlib.Path
        ``source_data/STab1``.

    Returns
    -------
    dict of pandas.DataFrame
        Keys ``ids``, ``ko``, ``cur``, ``new``, ``rc``, ``korc``.
    """
    sd = Path(sd)
    return {
        "ids": _read(sd / "per_acquisition_identifications.tsv"),
        "ko": _read(sd / "per_acquisition_KO_coverage.tsv"),
        "cur": _read(sd / "support_97_99_identification_comparison.tsv"),
        "new": _read(sd / "check_new_rows_95_90_75_50.tsv"),
        "rc": _read(sd / "independent_recount_identifications.tsv"),
        "korc": _read(sd / "independent_recount_KO.tsv"),
    }


def merged(d):
    """Join peptide and KO producer tables one-to-one on acquisition and arm.

    Parameters
    ----------
    d : dict
        Output of :func:`load`.

    Returns
    -------
    pandas.DataFrame
        105 rows, one per acquisition and arm.
    """
    t = d["ids"].merge(d["ko"][["sample", "arm", "compatible_KOs", "baseline_compatible_KOs", "baseline_KOs_retained",
                                "KO_lost", "KO_gained", "KO_recovery"]], on=["sample", "arm"], how="inner",
                       validate="1:1")
    assert len(t) == N_ROWS and set(t.arm) == set(ARM_ID) and t["sample"].nunique() == N_ACQUISITIONS
    return t


def checks(d):
    """Repeat the producer's checks.

    Checks: the earlier 45 rows and the 60 new rows equal the producer table under the
    column mapping; the two sets do not overlap; the independent recount agrees on all
    105 rows; all 90 clustered rows are below their unclustered baseline.

    Parameters
    ----------
    d : dict
        Output of :func:`load`.

    Returns
    -------
    (pandas.DataFrame, bool)
        The check table (same columns as the producer's CHECKS.tsv) and whether all pass.
    """
    mapped = d["ids"][["sample", "arm"] + list(MAP)].rename(columns=MAP)
    rows = []

    def cmp(name, ref, n_expected):
        m = ref.merge(mapped, on=["sample", "arm"], suffixes=("_ref", "_map"), how="left", indicator=True)
        missing = int((m["_merge"] != "both").sum())
        int_mis = sum(int((m[c + "_ref"] != m[c + "_map"]).sum()) for c in INTS)
        flt_exact = sum(int((m[c + "_ref"] != m[c + "_map"]).sum()) for c in FLOATS)
        flt_max = max(float((m[c + "_ref"] - m[c + "_map"]).abs().max()) for c in FLOATS)
        rows.append((name, len(ref), n_expected, missing, int_mis, flt_exact, f"{flt_max:.3g}"))
        return len(ref) == n_expected and missing == 0 and int_mis == 0 and flt_max < FLOAT_TOL

    ok = cmp("current_45_rows_vs_mapping", d["cur"], 45)
    ok &= cmp("new_60_rows_vs_mapping", d["new"], 60)
    ok &= set(map(tuple, d["cur"][["sample", "arm"]].values)).isdisjoint(map(tuple, d["new"][["sample", "arm"]].values))

    t = merged(d)
    r2 = t.merge(d["korc"], on=["sample", "arm"], suffixes=("", "_rc"))
    r1 = t.merge(d["rc"], on=["sample", "arm"], suffixes=("", "_rc"))
    rc_mis = int((r1.peptides != r1.peptides_rc).sum() + (r1.lost != r1.lost_rc).sum() + (r1.gained != r1.gained_rc).sum()
                 + (r1.protein_entries != r1.entries).sum() + (r2.compatible_KOs != r2.KOs).sum()
                 + ((r2.KO_recovery - r2.KO_recovery_rc).abs() > FLOAT_TOL).sum())
    rows.append(("independent_recount_vs_table", len(r1), N_ROWS, N_ROWS - len(r2), rc_mis, "", ""))
    ok &= len(r1) == N_ROWS and len(r2) == N_ROWS and rc_mis == 0

    med = stat_medians(d)
    for arm, row in med.iterrows():
        rows.append((f"median_{arm}", f"reduction {row.red:.3f}%", f"KO {100*row.ko:.2f}%", "", "", "", ""))
    below = stat_rows_below_baseline(d)
    rows.append(("clustered_rows_below_baseline", below, 90, "", "", "", ""))
    ok &= below == 90
    cols = ["check", "rows_or_value", "expected_or_value", "missing", "int_mismatches",
            "float_not_bit_identical", "float_max_abs_diff"]
    return pd.DataFrame(rows, columns=cols), bool(ok)


def build_table(d):
    """Build Supplementary Table 1 exactly as the producer did.

    Parameters
    ----------
    d : dict
        Output of :func:`load`.

    Returns
    -------
    pandas.DataFrame
        105 rows x 21 columns, sorted by acquisition then arm order.
    """
    t = merged(d)
    t["_a"] = t.arm.map(ARM_ORDER)
    t = t.sort_values(["sample", "_a"]).reset_index(drop=True)
    out = pd.DataFrame({
        "Acquisition": t["sample"],
        "Arm": t["arm"],
        "Linclust minimum sequence identity (%)": t["arm"].map(ARM_ID).astype("Int64"),
        "Searched protein entries (n)": t["protein_entries"],
        "Protein entry ratio (clustered/unclustered)": t["protein_entry_ratio"],
        "Accepted peptides (n)": t["peptides"],
        "Unclustered accepted peptides (n)": t["baseline_peptides"],
        "Peptide count ratio (clustered/unclustered)": t["peptide_count_ratio"],
        "Peptide reduction (%)": t["peptide_reduction_pct"],
        "Peptides lost vs unclustered (n)": t["lost"],
        "Peptides gained vs unclustered (n)": t["gained"],
        "Unclustered peptide identities retained (fraction)": t["baseline_identity_recovery"],
        "Reference-database peptides (n)": t["REF_peptides"],
        "Reference-database peptides recovered (n)": t["REF_recovered"],
        "Reference-database peptide recovery (fraction)": t["REF_recovery"],
        "Compatible KOs (n)": t["compatible_KOs"],
        "Unclustered compatible KOs (n)": t["baseline_compatible_KOs"],
        "Unclustered KOs retained (n)": t["baseline_KOs_retained"],
        "KOs lost vs unclustered (n)": t["KO_lost"],
        "KOs gained vs unclustered (n)": t["KO_gained"],
        "KO recovery (fraction)": t["KO_recovery"],
    })
    assert len(out) == N_ROWS
    return out


def write_table(out, out_dir):
    """Write the table as TSV (shortest round-trip float text) and as a one-sheet XLSX.

    Parameters
    ----------
    out : pandas.DataFrame
        Output of :func:`build_table`.
    out_dir : pathlib.Path
        Destination directory.

    Returns
    -------
    (pathlib.Path, pathlib.Path)
        Paths of the TSV and XLSX files.
    """
    out_dir = Path(out_dir)
    txt = out.copy()
    for c in txt.columns:   # shortest round-trip text for floats, as in the producer tables
        if txt[c].dtype == float:
            txt[c] = [repr(float(v)) for v in txt[c]]
    tsv, xlsx = out_dir / "Supplementary_Table_1.tsv", out_dir / "Supplementary_Table_1.xlsx"
    txt.to_csv(tsv, sep="\t", index=False)
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        out.to_excel(xw, sheet_name="Supplementary_Table_1", index=False)
        xw.sheets["Supplementary_Table_1"].freeze_panes = "C2"
    return tsv, xlsx


def compare_xlsx_values(new, expected):
    """Compare two workbooks cell by cell (values, sheet names, frozen panes).

    XLSX bytes contain a creation timestamp, so byte identity is not expected.

    Parameters
    ----------
    new, expected : pathlib.Path
        Workbooks to compare.

    Returns
    -------
    dict
        ``sheets``, ``cells_compared``, ``cells_differing``, ``freeze_panes`` and
        ``identical`` (bool).
    """
    import openpyxl
    a, b = openpyxl.load_workbook(new), openpyxl.load_workbook(expected)
    res = {"sheets": [a.sheetnames, b.sheetnames], "cells_compared": 0, "cells_differing": 0}
    same = a.sheetnames == b.sheetnames
    for name in a.sheetnames:
        wa, wb = a[name], b[name]
        same &= (wa.max_row, wa.max_column) == (wb.max_row, wb.max_column)
        res["freeze_panes"] = [wa.freeze_panes, wb.freeze_panes]
        same &= wa.freeze_panes == wb.freeze_panes
        for ra, rb in zip(wa.iter_rows(values_only=True), wb.iter_rows(values_only=True)):
            for x, y in zip(ra, rb):
                res["cells_compared"] += 1
                if x != y and not (x is None and y is None):
                    res["cells_differing"] += 1
    res["identical"] = bool(same and res["cells_differing"] == 0)
    return res


# ---------------------------------------------------------------------------------------
# Statistics printed in the manuscript
# ---------------------------------------------------------------------------------------
def stat_medians(d):
    """Median peptide reduction (%) and median KO recovery (fraction) per clustered arm.

    Statistical unit: acquisition; n = 15 acquisitions per arm, paired with the same
    acquisition's unclustered search. Descriptive median, no test.

    Parameters
    ----------
    d : dict
        Output of :func:`load`.

    Returns
    -------
    pandas.DataFrame
        Index arm; columns ``red`` (median peptide reduction, %), ``ko`` (median KO
        recovery, fraction) and ``n`` (acquisitions).
    """
    t = merged(d)
    return (t[t.arm != "unclustered"].groupby("arm")
            .agg(red=("peptide_reduction_pct", "median"), ko=("KO_recovery", "median"),
                 n=("sample", "nunique")))


def stat_ko_median_range_pct(d):
    """Lowest and highest per-arm median KO recovery, in percent.

    Statistical unit: acquisition; each median over n = 15 acquisitions, six arms.

    Returns
    -------
    (float, float)
        (minimum, maximum) over the six clustered arms of median KO recovery x 100.
    """
    k = stat_medians(d)["ko"] * 100
    return float(k.min()), float(k.max())


def stat_rows_below_baseline(d):
    """Number of clustered acquisition-arm rows with fewer accepted peptides than unclustered.

    Statistical unit: acquisition; 15 acquisitions x 6 arms = 90 rows.

    Returns
    -------
    int
    """
    t = merged(d)
    return int((t[t.arm != "unclustered"].peptide_count_ratio < 1).sum())


def stat_acquisitions_reduced_at_every_threshold(d):
    """Number of acquisitions whose accepted peptide count fell at all six thresholds.

    Statistical unit: acquisition; n = 15.

    Returns
    -------
    int
    """
    t = merged(d)
    c = t[t.arm != "unclustered"]
    per = c.groupby("sample").apply(lambda g: bool((g.peptide_count_ratio < 1).all()) and len(g) == 6,
                                    include_groups=False)
    return int(per.sum())


def stat_n_acquisitions(d):
    """Number of distinct CAMPI acquisitions in the table (statistical unit: acquisition)."""
    return int(merged(d)["sample"].nunique())


def stat_n_rows(d):
    """Number of acquisition-arm configurations in the table."""
    return int(len(build_table(d)))
